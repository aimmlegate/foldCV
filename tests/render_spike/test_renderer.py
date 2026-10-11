"""Real DOCX-to-PDF regression checks for the synthetic render spike."""

from __future__ import annotations

import json
import re
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from PIL import Image, ImageChops
from docx import Document
from docx.oxml.ns import qn

from spikes.render import renderer


ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "spikes/render/fixtures"
BASELINES = Path(__file__).resolve().parent / "baselines"
CASES = {
    "short": (1, ["Мария Примерова", "35%", "Лаборатория Север", "НАВЫКИ"]),
    "long-main": (4, ["Тестовая команда 01", "Тестовая команда 10", "20%", "НАВЫКИ"]),
    "long-sidebar": (3, ["RabbitMQ", "Учебная программа 12", "Синтетический курс 12", "English — B2"]),
}


def _normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


class RenderSpikeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        renderer.require_environment()
        cls.temp = tempfile.TemporaryDirectory(prefix="foldcv-render-tests-")
        cls.output = Path(cls.temp.name)
        for stem in CASES:
            data = json.loads((FIXTURES / f"{stem}.json").read_text(encoding="utf-8"))
            docx_path = cls.output / f"{stem}.docx"
            renderer.render_docx(data, docx_path)
            renderer.convert_pdf(docx_path, cls.output / f"{stem}.pdf")

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp.cleanup()

    def test_content_structure_fonts_and_geometry(self) -> None:
        for stem, (expected_pages, markers) in CASES.items():
            with self.subTest(fixture=stem):
                data = json.loads((FIXTURES / f"{stem}.json").read_text(encoding="utf-8"))
                docx_path = self.output / f"{stem}.docx"
                document = Document(docx_path)
                self.assertEqual(len(document.tables) - 1, expected_pages)
                self.assertEqual(len(document.inline_shapes), len(data["focus"]) + 1)
                xml_text = " ".join(node.text or "" for node in document.element.body.iter(qn("w:t")))
                pdf = pymupdf.open(self.output / f"{stem}.pdf")
                self.assertEqual(len(pdf), expected_pages)
                text = _normalized(" ".join(page.get_text() for page in pdf))
                for marker in markers:
                    self.assertIn(marker, xml_text)
                    self.assertIn(marker, text)
                expected = [data["candidate"][key] for key in ("name", "title", "summary")]
                for project in data["projects"]:
                    expected.extend(project[key] for key in ("company", "period", "stack", "description", "role"))
                    expected.extend(project["tasks"] + project["results"])
                expected.extend(item for section in data["sidebar"] for item in section["items"])
                expected.extend(data["achievements"])
                expected.extend(value for group in data["skill_groups"]
                                for value in (group["title"], group["items"]))
                for value in expected:
                    self.assertIn(_normalized(value), text)
                for number, page in enumerate(pdf, 1):
                    self.assertAlmostEqual(page.rect.width, 595.3, delta=0.5)
                    self.assertAlmostEqual(page.rect.height, 841.9, delta=0.5)
                    self.assertGreater(len(_normalized(page.get_text())), 250)
                    fonts = {name.split("+")[-1] for _, _, _, name, *_ in page.get_fonts(full=True)}
                    self.assertTrue(fonts <= {"SourceSansPro-Regular", "SourceSansPro-Bold"}, fonts)
                    for block in page.get_text("dict")["blocks"]:
                        for line in block.get("lines", []):
                            for span in line["spans"]:
                                box = pymupdf.Rect(span["bbox"])
                                self.assertGreaterEqual(box.x0, 0)
                                self.assertGreaterEqual(box.y0, 0)
                                self.assertLessEqual(box.x1, page.rect.width)
                                self.assertLessEqual(box.y1, page.rect.height)
                    self.assertIn("Synthetic local document", page.get_text())
                    self.assertTrue((BASELINES / f"{stem}-{number:02d}.png").is_file())

    def test_page_images_match_reviewed_baselines(self) -> None:
        for stem, (expected_pages, _) in CASES.items():
            pdf = pymupdf.open(self.output / f"{stem}.pdf")
            for number in range(1, expected_pages + 1):
                with self.subTest(fixture=stem, page=number):
                    pix = pdf[number - 1].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                    actual = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    with Image.open(BASELINES / f"{stem}-{number:02d}.png") as reference:
                        self.assertEqual(actual.size, reference.size)
                        diff = ImageChops.difference(actual, reference.convert("RGB"))
                        changed = sum(1 for value in diff.convert("L").get_flattened_data() if value > 8)
                        self.assertLessEqual(changed / (actual.width * actual.height), 0.002)

    def test_ordinary_edit_survives_reconversion_with_later_content(self) -> None:
        path = self.output / "long-main.docx"
        document = Document(path)
        changed = False
        for table in document.tables[1:]:
            for paragraph in table.cell(0, 1).paragraphs:
                if "Построил сервис обработки потока для этапа 1" in paragraph.text:
                    for run in paragraph.runs:
                        if "Построил сервис обработки потока" in run.text:
                            run.text = run.text.replace("Построил сервис", "Разработал сервис")
                            changed = True
                            break
        self.assertTrue(changed)
        edited = self.output / "edited.docx"
        document.save(edited)
        converted = self.output / "edited.pdf"
        renderer.convert_pdf(edited, converted)
        pdf = pymupdf.open(converted)
        text = _normalized(" ".join(page.get_text() for page in pdf))
        self.assertIn("Разработал сервис обработки потока", text)
        self.assertNotIn("Построил сервис обработки потока для этапа 1", text)
        self.assertIn("Тестовая команда 10", text)
        self.assertIn("НАВЫКИ", text)
        self.assertEqual(len(pdf), 4)

    def test_missing_dependencies_and_oversized_block_fail_visibly(self) -> None:
        with patch.object(renderer, "SOFFICE", Path("/missing/soffice")):
            with self.assertRaisesRegex(RuntimeError, "converter missing"):
                renderer.require_environment()
        with patch.object(renderer, "APP_FONTS", Path("/missing/fonts")):
            with self.assertRaisesRegex(RuntimeError, "font missing"):
                renderer.require_environment()
        data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
        data["candidate"]["summary"] = "Очень длинный текст " * 500
        with self.assertRaisesRegex(ValueError, "page budget"):
            renderer.render_docx(data, self.output / "oversized.docx")
        data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
        data["candidate"]["name"] = "Очень длинное демонстрационное имя " * 8
        with self.assertRaisesRegex(ValueError, "header line"):
            renderer.render_docx(data, self.output / "long-name.docx")

    def test_optional_sections_can_be_empty(self) -> None:
        data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
        data.update(focus=[], sidebar=[], projects=[], achievements=[], skill_groups=[])
        path = self.output / "minimal.docx"
        renderer.render_docx(data, path)
        renderer.convert_pdf(path, self.output / "minimal.pdf")
        pdf = pymupdf.open(self.output / "minimal.pdf")
        self.assertEqual(len(pdf), 1)
        self.assertIn(data["candidate"]["summary"], _normalized(pdf[0].get_text()))

    def test_empty_skill_groups_leave_no_heading(self) -> None:
        data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
        data["skill_groups"] = [
            {"title": "EMPTY OPTIONAL SKILL GROUP", "items": ""},
            {"title": "BLANK OPTIONAL SKILL GROUP", "items": " \n\t "},
            {"title": "VISIBLE SKILL GROUP", "items": "Python"},
        ]
        path = self.output / "empty-skills.docx"
        renderer.render_docx(data, path)
        renderer.convert_pdf(path, self.output / "empty-skills.pdf")
        document = Document(path)
        docx_text = " ".join(node.text or "" for node in document.element.body.iter(qn("w:t")))
        pdf = pymupdf.open(self.output / "empty-skills.pdf")
        pdf_text = _normalized(" ".join(page.get_text() for page in pdf))
        for text in (docx_text, pdf_text):
            self.assertNotIn("EMPTY OPTIONAL SKILL GROUP", text)
            self.assertNotIn("BLANK OPTIONAL SKILL GROUP", text)
            self.assertIn("VISIBLE SKILL GROUP", text)
        self.assertEqual(len(pdf), 1)

    def test_explicit_newlines_use_page_budget_and_preserve_short_text(self) -> None:
        data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
        data["candidate"]["summary"] = "\n".join(
            f"Synthetic line {number}" for number in range(1, 71))
        rejected = self.output / "seventy-lines.docx"
        with self.assertRaisesRegex(ValueError, "page budget"):
            renderer.render_docx(data, rejected)
        self.assertFalse(rejected.exists())

        lines = ["Synthetic line one", "Synthetic line two", "Synthetic line three"]
        data["candidate"]["summary"] = "\n".join(lines)
        accepted = self.output / "three-lines.docx"
        renderer.render_docx(data, accepted)
        renderer.convert_pdf(accepted, self.output / "three-lines.pdf")
        document = Document(accepted)
        docx_text = " ".join(node.text or "" for node in document.element.body.iter(qn("w:t")))
        pdf = pymupdf.open(self.output / "three-lines.pdf")
        self.assertEqual(len(pdf), 1)
        pdf_text = _normalized(pdf[0].get_text())
        for line in lines:
            self.assertIn(line, docx_text)
            self.assertIn(line, pdf_text)
        self.assertIn("Лаборатория Север", pdf_text)

    def test_conversion_timeout_fails_visibly(self) -> None:
        with patch.object(renderer.subprocess, "run", side_effect=subprocess.TimeoutExpired("soffice", 1)):
            with self.assertRaisesRegex(RuntimeError, "timed out"):
                renderer.convert_pdf(self.output / "short.docx", self.output / "timeout.pdf", timeout=1)


if __name__ == "__main__":
    unittest.main()
