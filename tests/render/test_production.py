"""Real local converter and bounded process regressions on synthetic resumes."""

from __future__ import annotations

import asyncio
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageChops
import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "services/api/src"))
from resume_api.lifecycle import OperationFiles
from resume_api.rendering import DocumentRenderer
from packages.resume_core.rendering.models import RenderDocument, RenderIdentity, RenderRequest, snapshot_digest
from packages.resume_core.rendering.template import PACKAGE_SHA256, RenderFailure, renderer_sha256, require_template
from packages.resume_core.rendering import template, worker

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "tests/fixtures/synthetic/render"
BASELINES = FIXTURES / "baselines"
CASES = {
    "short": (1, ["Мария Примерова", "35%", "Лаборатория Север", "НАВЫКИ"]),
    "long-main": (4, ["Тестовая команда 01", "Тестовая команда 10", "20%", "НАВЫКИ"]),
    "long-sidebar": (3, ["RabbitMQ", "Учебная программа 12", "Синтетический курс 12", "English — B2"]),
}


def request_for(data: dict, name: str) -> RenderRequest:
    document = RenderDocument.model_validate(data, by_name=True)
    identity = RenderIdentity.model_construct(id=name, project_revision=1, snapshot_sha256="0" * 64,
                                              template_id="generic", template_version="1.0.0",
                                              template_package_sha256=PACKAGE_SHA256, renderer_version="1.0.0",
                                              renderer_sha256=renderer_sha256())
    request = RenderRequest.model_construct(base_revision=1, snapshot=identity, document=document)
    request.snapshot.snapshot_sha256 = snapshot_digest(request)
    return RenderRequest.model_validate_json(request.model_dump_json(by_alias=True))


async def render(request: RenderRequest, tmp_path: Path) -> tuple[bytes, bytes, dict]:
    files = OperationFiles(tmp_path / "operations")
    with files.open():
        result = await DocumentRenderer().render(request, files)
    assert list((tmp_path / "operations").iterdir()) == []
    with zipfile.ZipFile(BytesIO(result)) as archive:
        assert archive.namelist() == ["resume.docx", "resume.pdf", "render.json"]
        return archive.read("resume.docx"), archive.read("resume.pdf"), json.loads(archive.read("render.json"))


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


@pytest.mark.slow
@pytest.mark.asyncio
async def test_real_renderer_content_layout_and_rasters(tmp_path):
    require_template()
    for stem, (pages, markers) in CASES.items():
        data = json.loads((FIXTURES / f"{stem}.json").read_text(encoding="utf-8"))
        request = request_for(data, stem)
        docx_bytes, pdf_bytes, meta = await render(request, tmp_path / stem)
        assert meta["snapshotSha256"] == request.snapshot.snapshot_sha256
        assert meta["templatePackageSha256"] == PACKAGE_SHA256
        assert meta["rendererSha256"] == renderer_sha256()
        assert meta["docxSha256"] == hashlib.sha256(docx_bytes).hexdigest()
        assert meta["pdfSha256"] == hashlib.sha256(pdf_bytes).hexdigest()
        assert meta["pageCount"] == pages
        assert meta["capabilities"] == {"editableDocx": True, "pdfDerivedFromDocx": True, "embeddedPdfFonts": True}
        with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
            assert b"<cp:coreProperties" in archive.read("docProps/core.xml")
            assert b"<dc:creator" not in archive.read("docProps/core.xml")
        doc = Document(BytesIO(docx_bytes))
        assert len(doc.inline_shapes) == len(data["focus"]) + 1
        assert len(doc.tables) - 1 == pages
        assert all(table.cell(0, 1).paragraphs for table in doc.tables[1:])
        word_text = normalized(" ".join(node.text or "" for node in doc.element.body.iter(qn("w:t"))))
        with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf:
            assert len(pdf) == pages
            assert not pdf.metadata.get("author") and not pdf.metadata.get("title")
            pdf_text = normalized(" ".join(page.get_text() for page in pdf))
            for value in markers:
                assert value in word_text
                assert value in pdf_text
            for value in worker._expected_text(request):
                assert normalized(value) in pdf_text
            for number, page in enumerate(pdf, 1):
                assert abs(page.rect.width - 595.3) < 1
                assert abs(page.rect.height - 841.9) < 1
                assert {name.split("+")[-1] for _, _, _, name, *_ in page.get_fonts(full=True)} <= {"SourceSansPro-Regular", "SourceSansPro-Bold"}
                pix = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                with Image.open(BASELINES / f"{stem}-{number:02d}.png") as baseline:
                    difference = ImageChops.difference(image, baseline.convert("RGB"))
                    changed = sum(1 for value in difference.convert("L").get_flattened_data() if value > 8)
                    assert changed / (image.width * image.height) <= 0.002


@pytest.mark.slow
@pytest.mark.asyncio
async def test_long_project_and_education_are_complete_and_editable(tmp_path):
    data = json.loads((FIXTURES / "long-sidebar.json").read_text(encoding="utf-8"))
    data["projects"][0]["tasks"] += [f"Synthetic extended project task {number:02d} with measurable delivery." for number in range(1, 30)]
    data["sidebar"][0]["items"] += [f"Synthetic education course {number:02d} at Example Institute." for number in range(1, 16)]
    request = request_for(data, "long-project-education")
    docx_bytes, pdf_bytes, meta = await render(request, tmp_path)
    assert meta["pageCount"] > 3
    document = Document(BytesIO(docx_bytes))
    word_text = normalized(" ".join(node.text or "" for node in document.element.body.iter(qn("w:t"))))
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf:
        pdf_text = normalized(" ".join(page.get_text() for page in pdf))
    for marker in ("Synthetic extended project task 01", "Synthetic extended project task 29",
                   "Synthetic education course 01", "Synthetic education course 15"):
        assert marker in word_text and marker in pdf_text
    target = next(paragraph for table in document.tables[1:] for paragraph in table.cell(0, 1).paragraphs
                  if "Synthetic extended project task 01" in paragraph.text)
    target.runs[-1].text = "Edited synthetic task in Word"
    edited = tmp_path / "edited.docx"
    document.save(edited)
    assert "Edited synthetic task in Word" in normalized(" ".join(node.text or "" for node in Document(edited).element.body.iter(qn("w:t"))))


@pytest.mark.slow
@pytest.mark.asyncio
async def test_oversize_individual_paragraphs_flow_across_pages(tmp_path):
    data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
    education = "EDU_BEGIN — " + ("Учебная программа по надежности сервисов. " * 42) + "EDU_MIDDLE — " + ("Практика распределенных систем и проверка данных. " * 42) + "EDU_END."
    description = "PROJECT_BEGIN — " + ("Платформа обработки событий для измеримых результатов. " * 34) + "PROJECT_MIDDLE — " + ("Проверка очередей, схем и восстановления после ошибок. " * 34) + "PROJECT_END."
    task = "TASK_BEGIN — " + ("Построила проверку потоков и журналов с точными метриками. " * 34) + "TASK_MIDDLE — " + ("Согласовала этапы запуска и восстановление после сбоев. " * 34) + "TASK_END."
    assert all(len(value) < 4000 for value in (education, description, task))
    data["sidebar"] = [{"title": "ОБРАЗОВАНИЕ", "items": [education]}]
    data["projects"][0]["description"] = description
    data["projects"][0]["tasks"] = [task]
    request = request_for(data, "flowing-paragraphs")
    docx_bytes, pdf_bytes, metadata = await render(request, tmp_path)
    document = Document(BytesIO(docx_bytes))
    word_text = normalized(" ".join(node.text or "" for node in document.element.body.iter(qn("w:t"))))
    assert "(continued)" not in word_text
    assert word_text.count("•") == 4
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf:
        assert len(pdf) == metadata["pageCount"] and len(pdf) >= 3
        page_text = [normalized(page.get_text()) for page in pdf]
        for text in page_text:
            body = text.replace(data["candidate"]["name"], "").replace(data["candidate"]["title"], "")
            assert normalized(body.replace(data["footer"], ""))
    for prefix in ("EDU", "PROJECT", "TASK"):
        locations = []
        for marker in (f"{prefix}_BEGIN", f"{prefix}_MIDDLE", f"{prefix}_END"):
            assert marker in word_text
            matches = [number for number, text in enumerate(page_text) if marker in text]
            assert len(matches) == 1
            locations.append(matches[0])
        assert locations == sorted(locations)
        assert len(set(locations)) >= 2


@pytest.mark.slow
@pytest.mark.asyncio
async def test_unknown_focus_icon_fails_with_a_public_code_and_cleans(tmp_path):
    data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
    data["focus"][0]["icon"] = "external"
    files = OperationFiles(tmp_path / "operations")
    with pytest.raises(RenderFailure) as caught:
        with files.open():
            await DocumentRenderer().render(request_for(data, "unknown-icon"), files)
    assert caught.value.code == "UNSUPPORTED_FOCUS_ICON"
    assert "Мария Примерова" not in str(caught.value)
    assert list((tmp_path / "operations").iterdir()) == []


@pytest.mark.slow
@pytest.mark.asyncio
async def test_empty_optional_sections_disappear(tmp_path):
    data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
    data["focus"] = []
    data["sidebar"] = [{"title": "EMPTY EDUCATION", "items": ["  "]}]
    data["projects"] = []
    data["achievements"] = [" "]
    data["skill_groups"] = [{"title": "EMPTY SKILLS", "items": " "}]
    docx_bytes, pdf_bytes, _ = await render(request_for(data, "minimal"), tmp_path)
    document = Document(BytesIO(docx_bytes))
    word_text = normalized(" ".join(node.text or "" for node in document.element.body.iter(qn("w:t"))))
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf:
        assert len(pdf) == 1
        pdf_text = normalized(pdf[0].get_text())
    for text in (word_text, pdf_text):
        assert data["candidate"]["summary"] in text
        assert "EMPTY EDUCATION" not in text
        assert "EMPTY SKILLS" not in text
        assert "ПРОЕКТЫ" not in text


def test_bounds_and_unavailable_assets_fail_visibly(tmp_path, monkeypatch):
    data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
    data["focus"] = data["focus"] * 3
    with pytest.raises(ValueError):
        request_for(data, "too-many-focus")
    data["focus"] = [{"label": "Bad", "duration": "", "icon": "external"}]
    assert request_for(data, "unknown-icon")
    manifest = template.require_docx_template()
    with pytest.raises(RenderFailure, match="UNSUPPORTED_FOCUS_ICON"):
        template.icon_path("external", manifest)
    monkeypatch.setattr(template, "SOFFICE", tmp_path / "missing")
    with pytest.raises(RenderFailure, match="CONVERTER_UNAVAILABLE"):
        template.require_template()
    wrong_converter = tmp_path / "soffice"
    wrong_converter.write_bytes(b"wrong binary")
    monkeypatch.setattr(template, "SOFFICE", wrong_converter)
    with pytest.raises(RenderFailure, match="CONVERTER_MISMATCH"):
        template.require_template()
    # A stand-in that matches the pin, so the font checks run without the installed converter.
    pinned = tmp_path / "pinned-soffice"
    pinned.write_bytes(b"pinned binary")
    monkeypatch.setattr(template, "require_docx_template", lambda: {
        **manifest, "converter": {"sha256": hashlib.sha256(b"pinned binary").hexdigest()}})
    monkeypatch.setattr(template, "SOFFICE", pinned)
    monkeypatch.setattr(template, "APP_FONTS", tmp_path / "missing-fonts")
    with pytest.raises(RenderFailure, match="FONT_UNAVAILABLE"):
        template.require_template()
    wrong_fonts = tmp_path / "wrong-fonts"
    wrong_fonts.mkdir()
    for name in ("SourceSansPro-Bold.ttf", "SourceSansPro-Regular.ttf"):
        (wrong_fonts / name).write_bytes(b"wrong font")
    monkeypatch.setattr(template, "APP_FONTS", wrong_fonts)
    with pytest.raises(RenderFailure, match="FONT_MISMATCH"):
        template.require_template()
    data["focus"] = []
    data["candidate"]["summary"] = "Very long indivisible text " * 300
    request = request_for(data, "unsupported-summary")
    with pytest.raises(ValueError, match="page budget"):
        worker.layout.render_docx(request.document.model_dump(), tmp_path / "unsupported.docx")
    assert not (tmp_path / "unsupported.docx").exists()

    class SlowProcess:
        killed = False
        calls = 0

        def wait(self, timeout):
            self.calls += 1
            if self.calls == 1:
                raise subprocess.TimeoutExpired("soffice", timeout)
            return 0

        def kill(self):
            self.killed = True

    slow = SlowProcess()
    monkeypatch.setattr(worker.subprocess, "Popen", lambda *args, **kwargs: slow)
    with pytest.raises(RenderFailure, match="CONVERTER_TIMEOUT"):
        worker._convert(tmp_path / "input.docx", tmp_path / "output.pdf", tmp_path)
    assert slow.killed

    large = tmp_path / "large.docx"
    large.write_bytes(os.urandom(17 * 1024 * 1024))
    small = tmp_path / "small.pdf"
    small.write_bytes(b"pdf")
    with pytest.raises(RenderFailure, match="RENDER_OUTPUT_LIMIT"):
        worker._write_zip(tmp_path / "large.zip", large, small, {})


@pytest.mark.asyncio
async def test_worker_cancel_stops_group_and_cleans_workspace(tmp_path, monkeypatch):
    script = tmp_path / "slow-worker.py"
    script.write_text("from pathlib import Path\nimport subprocess,sys,time\nsubprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])\nPath('started').write_text('yes')\ntime.sleep(60)\n")
    import resume_api.rendering as service_rendering
    monkeypatch.setattr(service_rendering, "WORKER", script)
    data = json.loads((FIXTURES / "short.json").read_text(encoding="utf-8"))
    request = request_for(data, "cancel")
    files = OperationFiles(tmp_path / "operations")

    async def run():
        with files.open():
            return await DocumentRenderer().render(request, files)

    task = asyncio.create_task(run())
    for _ in range(100):
        if files.path is not None and (files.path / "started").exists():
            break
        await asyncio.sleep(0.01)
    assert files.process is not None
    group_id = files.process.pid
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert list((tmp_path / "operations").iterdir()) == []
    from resume_api.cli.runner import group_members
    assert group_members(group_id) == []
