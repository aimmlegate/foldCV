"""Original evidence stays readable without opening a viewer or granting edits."""

import asyncio
from io import BytesIO
import json

import httpx
import pymupdf
import pytest
from docx import Document

from packages.resume_core.agent_store import run_in
from packages.resume_core.extraction.parser import extract
from packages.resume_core.gui_normalization import canonical_source
from packages.resume_core.models.agent import AgentQuery
from packages.resume_core.models.contracts import SourceOverlay
from tests.recovery.test_agent_tools import tools_setup


def original_file(kind):
    if kind == "pdf":
        with pymupdf.open() as document:
            document.new_page().insert_text((72, 72), "Original-only experience: 2018-2020")
            document.new_page().insert_text((72, 72), "Original-only skill: Python")
            return document.tobytes()
    if kind == "docx":
        document = Document()
        document.add_heading("Original-only experience: 2018-2020", 1)
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Original-only skill: Python"
        document.sections[0].header.paragraphs[0].text = "Original header"
        data = BytesIO(); document.save(data)
        return data.getvalue()
    prefix = "# " if kind == "md" else ""
    return (prefix + "Original-only experience: 2018-2020\r\n\r\n  Original-only skill: Python\n").encode()


def install_original(projects, workspace, kind="txt", *, data=None):
    data = original_file(kind) if data is None else data
    source = canonical_source(extract(data, f"Original.{kind}"))
    workspace.project.sources.append(source)
    projects.originals[(workspace.project_id, source.id)] = data
    return source, data


async def read_sources(session, ids=()):
    cursor, sources, fragments, pages = None, [], {}, 0
    while True:
        reply = await session.call("get_original_source", AgentQuery(kind="source", ids=list(ids), cursor=cursor))
        assert "error" not in reply, reply
        assert len(json.dumps(reply, ensure_ascii=False).encode()) < 16 * 1024 + 200
        payload = json.loads(reply["content"])
        if "fragment" in payload:
            prior = fragments.get(payload["itemIndex"], "")
            assert payload["offset"] == len(prior)
            fragments[payload["itemIndex"]] = prior + payload["fragment"]
            if payload["itemComplete"]:
                sources.extend(json.loads(fragments.pop(payload["itemIndex"]))["sources"])
        else:
            sources.extend(payload["sources"])
        pages += 1
        cursor = reply["snapshot"]["nextCursor"]
        if cursor is None:
            assert not fragments
            return sources, pages


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", ["pdf", "docx", "txt", "md"])
async def test_original_formats_keep_omitted_text_and_structure_separate_from_edits(tmp_path, kind):
    session, service, workspace, fetcher = tools_setup(tmp_path, allow_edits=False, allow_web=False)
    try:
        source, data = install_original(service.projects, workspace, kind)
        service.projects.catalog.update(workspace.project_id, read_only=True)
        source.units[0].disposition = "archived"
        source.overlays.append(SourceOverlay(id="overlay_1", revision=1, unitId=source.units[0].id,
            text="Edited overlay must not replace the original", origin="user"))
        next(block for block in workspace.project.document.blocks() if block.type == "paragraph").content[0].text = "Edited layout"
        before = workspace.model_dump_json()
        _, run, journal = run_in(service.store.load(workspace.project_id), session.run.id)
        prompt = session.book.initial(run, journal)
        access = json.loads(prompt.split("\n", 1)[1])["originalSourceAccess"]
        assert access["tool"] == "get_original_source" and access["sourceCount"] == 1
        assert access["sources"][0]["name"] == f"Original.{kind}"
        assert access["sources"][0]["extractedTextStatus"] == "available"
        sources, _ = await read_sources(session)
        units = [unit for part in sources for unit in part["units"]]
        assert units == [unit.model_dump(mode="json", by_alias=True) for unit in source.units]
        text = "".join(unit["text"] for unit in units)
        assert "Original-only experience: 2018-2020" in text and "Original-only skill: Python" in text
        assert "Edited overlay" not in text and "Edited layout" not in text
        assert sources[0]["mimeType"] == source.mime_type and sources[0]["sha256"] == source.byte_sha256
        assert all(unit["structure"] for unit in units)
        if kind == "pdf":
            assert {unit["location"]["index"] for unit in units} == {0, 1}
        if kind == "docx":
            assert any(unit["location"]["kind"] == "table_cell" for unit in units)
            assert "Original header" in text
        selected, _ = await read_sources(session, [source.units[0].id])
        assert [unit["id"] for part in selected for unit in part["units"]] == [source.units[0].id]
        denied = await session.call("get_original_source", AgentQuery(kind="source", ids=["foreign_source"]))
        assert denied["error"]["code"] == "AGENT_NOT_FOUND"
        assert not fetcher.calls and workspace.model_dump_json() == before
        assert service.projects.originals[(workspace.project_id, source.id)] == data
        assert (await session.call("apply_edit_batch", {}))["error"]["code"] == "AGENT_TOOL_DENIED"
    finally:
        await session.close()


@pytest.mark.asyncio
async def test_empty_and_unavailable_extraction_are_listed_with_page_warnings(tmp_path):
    session, service, workspace, _ = tools_setup(tmp_path, allow_edits=False)
    try:
        with pymupdf.open() as document:
            document.new_page()
            empty, _ = install_original(service.projects, workspace, "pdf", data=document.tobytes())
        unavailable, _ = install_original(service.projects, workspace)
        unavailable.units = []
        sources, _ = await read_sources(session)
        by_id = {part["id"]: part for part in sources}
        assert by_id[empty.id]["extractedTextStatus"] == "empty"
        assert by_id[empty.id]["warningDetails"] == [{"code": "NO_TEXT", "page": 1}]
        assert by_id[unavailable.id]["extractedTextStatus"] == "unavailable"
        assert by_id[unavailable.id]["units"] == []
    finally:
        await session.close()


@pytest.mark.asyncio
async def test_long_original_is_complete_across_pages_and_fragments_and_keeps_stale_guard(tmp_path):
    session, service, workspace, _ = tools_setup(tmp_path, allow_edits=False)
    try:
        expected = [f"ROW_{i:03}: " + "Ж😀" * 120 + "\n" for i in range(150)]
        expected.insert(70, "OVERSIZED: " + '"\\😀' * 9000 + "\n")
        source, _ = install_original(service.projects, workspace, data="".join(expected).encode())
        source.units.reverse()
        sources, pages = await read_sources(session, [source.id])
        units = [unit for part in sources for unit in part["units"]]
        assert pages > 2 and [unit["text"] for unit in units] == expected
        assert len({unit["id"] for unit in units}) == len(expected)
        assert session.book.validate(workspace)
        source.units[0].text += " Changed original evidence"
        assert not session.book.validate(workspace)
    finally:
        await session.close()


@pytest.mark.asyncio
async def test_mcp_default_read_survives_reopen_and_portable_transfer(tmp_path):
    from packages.resume_core.rules import RulesStore
    from resume_api.rich_projects import RichProjectService
    from resume_api.agent.context import ReadBook, canonical_capture

    session, service, workspace, fetcher = tools_setup(tmp_path, allow_edits=False, allow_web=False)
    try:
        source, data = install_original(service.projects, workspace)
        assert service.projects.persist(workspace.project_id) == "durable"
        reopened = RichProjectService(tmp_path / "reopened", RulesStore(service.projects.storage.root), lambda: None)
        service.projects = reopened
        session.book = ReadBook(service, session.control)
        session.application = session.sdk.streamable_http_app()
        headers = {"Authorization": "Bearer " + session.token, "Accept": "application/json, text/event-stream"}
        async with session.application.router.lifespan_context(session.application):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=session), base_url="http://" + session.host, headers=headers) as client:
                response = await client.post("/mcp", json={"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                    "params": {"name": "get_original_source", "arguments": {}}})
                reply = json.loads(response.json()["result"]["content"][0]["text"])
                assert "Original-only experience" in reply["content"] and "error" not in reply
        copied = reopened.import_portable(reopened.portable(workspace.project_id))
        capture = canonical_capture(reopened.read(copied), AgentQuery(kind="source"))
        assert capture.payload["sources"][0]["units"] == source.units
        assert reopened.originals[(copied, source.id)] == data and not fetcher.calls
    finally:
        await session.close()


@pytest.mark.asyncio
async def test_queued_read_only_assistant_reads_original_without_viewer_or_web(tmp_path):
    from packages.resume_core.background_supervisor import ExecutionOutcome
    from tests.recovery.test_background_agent import environment

    service, factory, workspace, payload, _connection, _old = await environment(tmp_path)
    current = service.projects.workspaces[workspace.project_id]
    source, data = install_original(service.projects, current, "md")
    assert service.projects.persist(workspace.project_id) == "durable"
    current = service.projects.read(workspace.project_id)
    payload.request.project_revision = current.project_revision
    payload.request.document_revision = current.document_revision
    payload.request.facts_revision = current.project.facts.revision
    payload.request.allow_edits = payload.request.allow_web = False
    checked = []

    async def execute(control, _actor):
        session = service.tools.session(control)
        try:
            assert "get_original_source" in session.names
            access = json.loads(session.prompt.split("\n", 1)[1])["originalSourceAccess"]
            assert access["sources"][0]["id"] == source.id
            parts, _ = await read_sources(session)
            checked.extend(unit["text"] for part in parts for unit in part["units"])
            assert "stored_extracted_text" == parts[0]["content"]
            return ExecutionOutcome(finished_at=service.supervisor.clock.now())
        finally:
            await session.close()

    factory.callbacks[payload.request.client_request_id] = execute
    try:
        job, _ = await service.submit(workspace.project_id, payload)
        await factory.actor(job)
        await asyncio.wait_for(service.supervisor.controls[job.id].task, 10)
        failure = service.supervisor.controls[job.id].run_task.exception()
        if failure is not None:
            raise failure
        assert "Original-only experience: 2018-2020" in "".join(checked)
        assert service.projects.originals[(workspace.project_id, source.id)] == data
        assert service.projects.storage.read_committed(workspace.project_id)[0] == current
    finally:
        await service.shutdown()
        factory.verify(); service.supervisor.clock.verify()
