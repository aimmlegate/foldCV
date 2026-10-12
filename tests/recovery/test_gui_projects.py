"""Connected creation, retained stages and protected in-memory project routes."""

import asyncio
import hashlib
from io import BytesIO
from pathlib import Path
import sys
import threading
from zipfile import ZipFile

import httpx
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "services/api/src"))

from packages.resume_core.rules import RulesStore
from packages.resume_core.rendering.docx_errors import failure
from resume_api.app import create_app
from resume_api.gui_projects import GuiProjectService, ProjectFailure
from resume_api.gui_admission import AdmissionCancelled, QueuedProvider
from resume_api.gui_content import generate_admitted
from resume_api.lifecycle import OperationFiles
from resume_api.providers.codex_cli import CodexCliProvider, _GENERATION_LOCK
from resume_api.providers.content import ContentFailure
from tests.recovery.generation_fixtures import CreationProvider, frozen_now
from tests.recovery.candidate_fixtures import ROOT
from tests.recovery.test_gui_content import request as content_request


def service(tmp_path, provider=None, renderer=None):
    provider = provider or CreationProvider()
    return GuiProjectService(tmp_path / "operations", RulesStore(tmp_path / "FoldCV"), lambda: (provider, "fixture"),
        clock=frozen_now, renderer=renderer)


async def finish(store, id, gid):
    await asyncio.wait_for(store.tasks[(id, gid)], timeout=40)
    return store.read(id)


@pytest.mark.asyncio
@pytest.mark.parametrize("example", ["E1", "E2"])
async def test_connected_creation_promotes_ids_and_real_docx_with_two_calls(example, tmp_path, monkeypatch):
    from packages.resume_core import dates, drafting
    from packages.resume_core.rendering import preparation
    def forbidden(*_args, **_kwargs):
        raise AssertionError("legacy editorial path")
    monkeypatch.setattr(dates, "old_job", forbidden)
    monkeypatch.setattr(drafting, "_recent_jobs", forbidden)
    monkeypatch.setattr(preparation, "prepare_document", forbidden)
    provider = CreationProvider(example)
    store = service(tmp_path, provider)
    data = (ROOT / f"{example}.txt").read_bytes()
    selected = await store.select(data, "pasted.txt", True)
    assert provider.calls == [] and (selected.project_revision, selected.document_revision) == (1, 1)
    source_id = selected.pre_draft.selected_source_id
    gid = store.start(selected.project_id)
    workspace = await finish(store, selected.project_id, gid)
    generation = workspace.generations[-1]
    assert generation.status == "completed", [(s.kind, s.error_code) for s in generation.stages]
    assert [stage.status for stage in generation.stages] == ["completed"] * 5
    assert all(stage.dependency_fingerprint == generation.dependencies.dependency_fingerprint for stage in generation.stages)
    assert workspace.pre_draft is None and (workspace.project_revision, workspace.document_revision) == (2, 2)
    assert workspace.project.sources[0].id == source_id
    assert store.originals[(workspace.project_id, source_id)] == data
    assert len(provider.calls) == len(workspace.project.operations) == 2
    assert [operation.usage.logical_attempts for operation in workspace.project.operations] == [1, 1]
    assert all(operation.dependencies.document_revision == 1 for operation in workspace.project.operations)
    assert all(operation.id in [s.operation_id for s in generation.stages] for operation in workspace.project.operations)
    assert all(not view.state.locked and view.previous is None for view in workspace.project.views)
    assert workspace.derived_state.document_revision == 2
    snapshot = workspace.project.renders[0]
    content = store.artifacts[(workspace.project_id, snapshot.id)]
    assert snapshot.docx.sha256 == hashlib.sha256(content).hexdigest() and snapshot.docx.size == len(content)
    assert snapshot.pdf is None and snapshot.project_revision == snapshot.document_revision == 2
    with ZipFile(BytesIO(content)) as archive:
        xml = archive.read("word/document.xml").decode()
        assert "Ирина Примерова" in xml if example == "E1" else "Алексей" in xml
        assert "mail@example.com" not in xml
        assert "Роль:" in xml if example == "E1" else "Role:" in xml
        if example == "E1":
            assert "w:br" in xml
    assert not list((tmp_path / "operations").glob("operation-*"))


@pytest.mark.asyncio
async def test_failed_draft_retry_reuses_normalization_without_extra_attempt(tmp_path):
    provider = CreationProvider(fail_kind="GuiDraftProposal")
    store = service(tmp_path, provider)
    selected = await store.select((ROOT / "E1.txt").read_bytes(), "source.txt", False)
    id = selected.project_id
    failed = await finish(store, id, store.start(id))
    assert failed.generations[-1].status == "failed"
    assert len(failed.pre_draft.facts) > 0 and len(provider.calls) == 2
    assert [op.usage.logical_attempts for op in failed.pre_draft.operations] == [1, 1]
    ids = [fact.id for fact in failed.pre_draft.facts]
    normalization_id = failed.pre_draft.operations[0].id
    provider.fail_kind = None
    result = await finish(store, id, store.start(id, failed.generations[-1].id))
    assert result.generations[-1].status == "completed" and len(provider.calls) == 3
    assert [fact.id for fact in result.project.facts] == ids
    assert result.generations[-1].stages[1].operation_id == normalization_id


@pytest.mark.asyncio
async def test_render_failure_keeps_document_and_retry_uses_no_model_calls(tmp_path):
    class BrokenRenderer:
        async def render(self, *_args):
            raise failure("DOCX_BLOCK_PAGE_BUDGET")
    provider = CreationProvider()
    store = service(tmp_path, provider, BrokenRenderer())
    selected = await store.select((ROOT / "E1.txt").read_bytes(), "pasted.txt", True)
    id = selected.project_id
    failed = await finish(store, id, store.start(id))
    assert failed.project is not None and failed.pre_draft is None
    assert failed.generations[-1].stages[-1].docx_error.code == "DOCX_BLOCK_PAGE_BUDGET"
    from resume_api.gui_rendering import GuiDocxRenderer
    store.renderer = GuiDocxRenderer()
    result = await finish(store, id, store.start(id, failed.generations[-1].id))
    assert result.generations[-1].status == "completed" and len(provider.calls) == 2
    assert (result.project_revision, result.document_revision) == (2, 2)


@pytest.mark.asyncio
async def test_empty_scan_failed_file_never_falls_back_and_replacements_keep_project(tmp_path):
    import pymupdf
    provider = CreationProvider()
    store = service(tmp_path, provider)
    selected = await store.select((ROOT / "E1.txt").read_bytes(), "pasted.txt", True)
    id = selected.project_id
    with pymupdf.open() as pdf:
        pdf.new_page()
        scanned = pdf.tobytes()
    for data, name in [(b" ", "empty.txt"), (scanned, "scan.pdf"), (b"broken", "broken.pdf")]:
        replacement = await store.select(data, name, False, id)
        assert replacement.pre_draft.selected_input == "file" and replacement.pre_draft.extracted_text.strip() == ""
        assert "NO_USABLE_TEXT" in replacement.pre_draft.warnings
        with pytest.raises(ProjectFailure, match="NO_USABLE_TEXT"):
            store.start(id)
    assert provider.calls == [] and len(store.originals) == len(replacement.pre_draft.sources) == 4
    replacement = await store.select((ROOT / "E1.txt").read_bytes(), "pasted.txt", True, id)
    result = await finish(store, id, store.start(id))
    assert result.project_id == id and result.generations[-1].source_ids == [replacement.pre_draft.selected_source_id]
    sent_ids = [item["id"] for item in provider.calls[0]["input"]["sources"]]
    assert sent_ids == result.generations[-1].source_ids
    assert len(result.project.sources) == 5


@pytest.mark.asyncio
async def test_cancelled_queued_content_has_zero_attempts_and_retained_extraction(tmp_path):
    entered = threading.Event()
    class Waiting(CreationProvider):
        def generate_queued(self, request, files, cancel, admitted):
            entered.set()
            while not cancel.wait(0.01):
                pass
            raise AdmissionCancelled()
    provider = Waiting()
    store = service(tmp_path, provider)
    selected = await store.select((ROOT / "E1.txt").read_bytes(), "source.txt", False)
    id, gid = selected.project_id, store.start(selected.project_id)
    assert await asyncio.to_thread(entered.wait, 2)
    before = store.read(id)
    assert before.pre_draft.operations[0].status == "queued"
    await store.cancel(id, gid)
    result = store.read(id)
    assert result.generations[-1].status == "cancelled"
    assert result.pre_draft.operations[0].usage.logical_attempts == 0
    assert result.pre_draft.operations[0].events[-1].type == "cancelled"
    assert result.pre_draft.extracted_text and provider.calls == []


def test_provider_queue_waits_shared_legacy_lock_and_busy_never_becomes_attempt(tmp_path, monkeypatch):
    provider = CodexCliProvider(Path("/nonexistent/synthetic"))
    dispatched = []
    monkeypatch.setattr(provider, "_generate", lambda *_args: dispatched.append(True))
    cancel = threading.Event()
    _GENERATION_LOCK.acquire()
    thread = threading.Thread(target=lambda: provider.generate_queued(content_request(), OperationFiles(tmp_path), cancel, lambda: dispatched.append("admitted")))
    thread.start()
    try:
        assert not dispatched
    finally:
        _GENERATION_LOCK.release()
    thread.join(2)
    assert not thread.is_alive() and dispatched == ["admitted", True]
    class InitiallyBusy:
        calls = 0
        def generate(self, *_args):
            self.calls += 1
            raise ContentFailure("BUSY" if self.calls == 1 else "MODEL_DENIED")
    provider = InitiallyBusy()
    with pytest.raises(ContentFailure) as error:
        generate_admitted(QueuedProvider(provider, lambda: None), content_request(), OperationFiles(tmp_path), cancel)
    assert error.value.code == "MODEL_DENIED" and error.value.logical_attempts == 1


@pytest.mark.asyncio
async def test_project_routes_retain_security_typed_errors_and_no_legacy_cap(tmp_path):
    origin, secret = "http://127.0.0.1:39329", "synthetic-project-secret-0123456789"
    dist = tmp_path / "dist"
    dist.mkdir()
    provider = CreationProvider()
    app = create_app(origin=origin, secret=secret, temp_root=tmp_path / "operations", dist=dist,
        data_root=tmp_path / "FoldCV", gui_provider=provider, gui_clock=frozen_now)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url=origin) as client:
        assert (await client.post("/v1/projects", headers={"Origin": origin}, json={"text": ""})).status_code == 401
        await client.post("/v1/session/exchange", headers={"Origin": origin}, json={"secret": secret})
        assert (await client.post("/v1/projects", json={"text": ""})).status_code == 403
        assert (await client.post("/v1/projects", headers={"Origin": origin}, content=b"x")).status_code == 415
        assert (await client.post("/v1/projects", headers={"Origin": origin, "Content-Type": "application/json"}, content=b'{"text":"a","text":"b"}')).status_code == 422
        response = await client.post("/v1/projects", headers={"Origin": origin}, json={"text": (ROOT / "E1.txt").read_text()})
        assert response.status_code == 202, response.text
        raw = response.json()
        assert raw["durability"] == "durable"
        assert (tmp_path / "FoldCV/projects" / raw["workspace"]["projectId"] / "workspace.json").is_file()
        id, gid = raw["workspace"]["projectId"], raw["generationId"]
        store = app.app.state.gui_projects
        workspace = await finish(store, id, gid)
        assert (await client.get(f"/v1/projects/{id}?private=content")).status_code == 403
        assert (await client.get(f"/v1/projects/{id}", headers={"Host": "localhost:39329"})).status_code == 403
        op = workspace.project.operations[0]
        assert (await client.get(f"/v1/projects/{id}/operations/{op.id}")).json()["usage"]["logicalAttempts"] == 1
        rid = workspace.project.renders[0].id
        download = await client.get(f"/v1/projects/{id}/renders/{rid}/docx")
        assert download.content == store.artifacts[(id, rid)] and download.headers["cache-control"] == "no-store"
        # Project batches retain their own counts beyond the legacy lifetime cap.
        for _ in range(21):
            result = await client.post("/v1/projects", headers={"Origin": origin}, json={"text": (ROOT / "E1.txt").read_text()})
            assert result.status_code == 202
            raw = result.json()
            completed = await finish(store, raw["workspace"]["projectId"], raw["generationId"])
            assert completed.generations[-1].status == "completed"
        assert len(provider.calls) == 44


@pytest.mark.asyncio
async def test_partial_extraction_retains_original_warnings_and_runs_only_selected_source(tmp_path):
    from docx import Document
    document = Document()
    for line in (ROOT / "E1.txt").read_text().splitlines():
        document.add_paragraph(line)
    fixture = Path(__file__).resolve().parents[1] / "fixtures/synthetic/import/rich.docx"
    with ZipFile(fixture) as archive:
        image = next(name for name in archive.namelist() if name.startswith("word/media/"))
        document.add_picture(BytesIO(archive.read(image)))
    buffer = BytesIO()
    document.save(buffer)
    provider = CreationProvider()
    store = service(tmp_path, provider)
    selected = await store.select(buffer.getvalue(), "partial.docx", False)
    assert "IMAGE" in selected.pre_draft.warnings and selected.pre_draft.extracted_text.strip()
    assert store.originals[(selected.project_id, selected.pre_draft.selected_source_id)] == buffer.getvalue()
    completed = await finish(store, selected.project_id, store.start(selected.project_id))
    assert completed.generations[-1].status == "completed" and len(provider.calls) == 2
    assert "IMAGE" in completed.project.sources[0].warnings


@pytest.mark.asyncio
async def test_changed_rules_invalidate_completed_retry_stages(tmp_path):
    provider = CreationProvider(fail_kind="GuiDraftProposal")
    store = service(tmp_path, provider)
    selected = await store.select((ROOT / "E1.txt").read_bytes(), "pasted.txt", True)
    failed = await finish(store, selected.project_id, store.start(selected.project_id))
    before = store.rules.load()
    store.rules.save(before.markdown.replace("15 years", "20 years"), before.revision)
    provider.fail_kind = None
    completed = await finish(store, selected.project_id, store.start(selected.project_id, failed.generations[-1].id))
    assert completed.generations[-1].status == "completed" and len(provider.calls) == 4
    assert completed.rules_used.revision == 2
    assert completed.generations[-1].stages[1].operation_id != failed.generations[-1].stages[1].operation_id


@pytest.mark.asyncio
async def test_host_detach_cancels_gui_queue_and_retains_project(tmp_path):
    entered = threading.Event()
    class Waiting(CreationProvider):
        def generate_queued(self, request, files, cancel, admitted):
            entered.set()
            cancel.wait(3)
            raise AdmissionCancelled()
    origin, secret = "http://127.0.0.1:39329", "synthetic-detach-secret-0123456789"
    dist = tmp_path / "dist"
    dist.mkdir()
    app = create_app(origin=origin, secret=secret, temp_root=tmp_path / "operations", dist=dist,
        data_root=tmp_path / "FoldCV", gui_provider=Waiting(), gui_clock=frozen_now)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url=origin) as client:
        headers = {"Origin": origin}
        await client.post("/v1/session/exchange", headers=headers, json={"secret": secret})
        response = await client.post("/v1/projects", headers=headers, json={"text": (ROOT / "E1.txt").read_text()})
        id = response.json()["workspace"]["projectId"]
        assert await asyncio.to_thread(entered.wait, 2)
        detached = await client.post("/v1/host/detach", headers=headers, json={})
        assert detached.status_code == 200 and detached.json()["state"] == "detached"
        workspace = (await client.get(f"/v1/projects/{id}")).json()["workspace"]
        assert workspace["generations"][-1]["status"] == "cancelled"
        assert workspace["preDraft"]["extractedText"] and workspace["preDraft"]["operations"][0]["usage"]["logicalAttempts"] == 0
