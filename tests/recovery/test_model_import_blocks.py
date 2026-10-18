"""Semantic source groups survive import without geometry or text-equality rules."""

from copy import deepcopy
from io import BytesIO
import threading
from zipfile import ZipFile

from lxml import etree
import pytest
from pydantic import ValidationError

from packages.resume_core.import_agent import ImportFailure, assemble_agent_import
from packages.resume_core.import_content import draft_content, update_content
from packages.resume_core.import_lists import format_import_lists
from packages.resume_core.models.importing import ImportContent, ImportContentFinish, ImportRead
from packages.resume_core.rendering.rich_docx import render_rich_docx
from packages.resume_core.rich_archive import open_rich_archive
from packages.resume_core.rich_presentation import authored_text
from packages.resume_core.rules import RulesStore
from resume_api.importer.session import ImportCoordinator, ImportSession
from resume_api.importer.result import import_workspace
from resume_api.rich_projects import RichProjectService
from tests.recovery.test_agent_import import draft_fixture, selected_service
from tests.recovery.test_import_content import paragraphs_fixture
from tests.recovery.test_import_lists import fixture as list_fixture


@pytest.mark.asyncio
async def test_semantic_blocks_join_split_period_and_stack_but_keep_role_and_distinct_contacts(tmp_path, monkeypatch):
    service, identifier, _runtime, _sessions = await selected_service(tmp_path, monkeypatch)
    texts = ["Candidate\n", "Tech Lead\n", "Company | 2012 –\n", "2020\n", "•\n", "• php, nodejs, nats,\n",
        "couchDB\n", "Phone: 123\n", "Phone: 124\n", "Work\n", "Work\n"]
    sources, parts, paragraphs = paragraphs_fixture(texts, pages=[1, 1, 1, 1, 2, 2, 2, 2, 2, 2, 3])
    original = deepcopy(sources)
    workspace = service.workspaces[identifier]
    workspace.pre_draft.sources = sources
    workspace.pre_draft.selected_source_id = sources[0].id
    draft = update_content(draft_fixture(), ImportContent(
        profile={"name": "Candidate", "refs": ["p1"]},
        facts=[{"kind": "employment", "employer": "Company", "role": "Tech Lead",
            "periods": [{"start": "2012", "end": "2020"}], "technologies": ["php", "nodejs", "nats", "couchDB"],
            "refs": ["p2-p7"]}],
        sections=[{"category": "identity", "blocks": [{"type": "heading", "refs": ["p1"]}]},
            {"category": "experience", "blocks": [{"type": "section", "refs": ["p10"]},
                {"type": "heading", "refs": ["p3-p4"]}, {"type": "paragraph", "refs": ["p2"]},
                {"type": "list_item", "refs": ["p5-p7"]},
                {"type": "duplicate", "refs": ["p11"], "duplicateOf": ["p10"]}]},
            {"category": "contacts", "blocks": [{"type": "paragraph", "refs": ["p8"]},
                {"type": "paragraph", "refs": ["p9"]}]}]), paragraphs, parts)
    draft.finished = True
    project = assemble_agent_import(workspace, draft)
    blocks = list(project.document.blocks())
    heading = next(b for b in blocks if authored_text(b) == "Company | 2012 – 2020")
    assert heading.type == "heading"
    assert any(b.type == "paragraph" and authored_text(b) == "Tech Lead" for b in blocks)
    assert any(b.type == "bullet_list" and authored_text(b) == "php, nodejs, nats, couchDB" for b in blocks)
    assert not any(authored_text(b) == "2020" or authored_text(b) == "couchDB" for b in blocks)
    assert {authored_text(b) for b in blocks if authored_text(b).startswith("Phone:")} == {"Phone: 123", "Phone: 124"}
    assert sum(authored_text(b) == "Work" for b in blocks) == 1
    excluded = next(row for row in project.import_coverage if row.unit_id == sources[0].units[10].id)
    assert excluded.disposition == "excluded" and "Model-marked duplicate" in excluded.reason
    assert sources[0].units[9].id in excluded.reason and not excluded.block_ids
    assert project.sources == original
    assert any(link.kind == "field_binding" and link.field_key.endswith(":end") for link in heading.provenance)
    with ZipFile(BytesIO(render_rich_docx(project.document, facts=project.facts))) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    docx_paragraphs = ["".join(p.xpath(".//w:t/text()", namespaces=ns)) for p in xml.xpath("//w:p", namespaces=ns)]
    assert "Company | 2012 – 2020" in docx_paragraphs
    assert "php, nodejs, nats, couchDB" in docx_paragraphs
    assert "2020" not in docx_paragraphs and "couchDB" not in docx_paragraphs
    await service.shutdown()


def test_semantic_blocks_reread_preserves_ref_order_duplicate_links_and_independent_saves():
    _sources, parts, paragraphs = paragraphs_fixture(["Candidate", "Tech Lead", "Company", "2012 –", "2020", "Candidate"])
    request = ImportContent(sections=[{"category": "identity", "blocks": [{"type": "heading", "refs": ["p1"]}]},
        {"category": "experience", "heading": "Work", "blocks": [
            {"type": "heading", "refs": ["p3", "p4-p5"]}, {"type": "paragraph", "refs": ["p2"]},
            {"type": "paragraph", "refs": ["p5", "p4"]},
            {"type": "duplicate", "refs": ["p6"], "duplicateOf": ["p1"]}]}])
    saved = update_content(draft_fixture(), request, paragraphs, parts)
    readable = [item["section"] for item in draft_content(saved, paragraphs) if "section" in item]
    assert readable[1]["blocks"][2]["refs"] == ["p5", "p4"]
    assert readable[1]["blocks"][3]["duplicateOf"] == ["p1"]
    assert update_content(saved, ImportContent(sections=readable), paragraphs, parts) == saved
    assert update_content(saved, request, paragraphs, parts) == saved
    changed = update_content(saved, ImportContent(facts=[{"text": "Dates need review."}]), paragraphs, parts)
    assert changed.blocks == saved.blocks
    cleared = update_content(changed, ImportContent(sections=[{"refs": ["p1-p6"]}]), paragraphs, parts)
    assert not any(block.model_block or block.duplicate_of for block in cleared.blocks)
    assert cleared.facts == changed.facts


@pytest.mark.parametrize("refs,target,visible", [(["p3"], ["p99"], ["p1"]),
    (["p3"], ["p3"], ["p3"]), (["p3"], ["p2"], ["p1"])])
def test_invalid_or_hidden_duplicate_targets_leave_saved_content_unchanged(refs, target, visible):
    _sources, parts, paragraphs = paragraphs_fixture(["Original", "Not placed", "Repeated"])
    saved = update_content(draft_fixture(), ImportContent(facts=[{"text": "Keep this fact."}]), paragraphs, parts)
    before = saved.model_dump_json()
    with pytest.raises(ImportFailure, match="IMPORT_REFERENCE_INVALID"):
        update_content(saved, ImportContent(sections=[{"blocks": [
            {"type": "paragraph", "refs": visible}, {"type": "duplicate", "refs": refs, "duplicateOf": target}]}]),
            paragraphs, parts)
    assert saved.model_dump_json() == before


def test_duplicate_chain_is_invalid_without_a_visible_target():
    _sources, parts, paragraphs = paragraphs_fixture(["Original", "Repeated", "Repeated again"])
    with pytest.raises(ImportFailure, match="IMPORT_REFERENCE_INVALID"):
        update_content(draft_fixture(), ImportContent(sections=[{"blocks": [
            {"type": "paragraph", "refs": ["p1"]},
            {"type": "duplicate", "refs": ["p2"], "duplicateOf": ["p3"]},
            {"type": "duplicate", "refs": ["p3"], "duplicateOf": ["p1"]}]}]), paragraphs, parts)


@pytest.mark.parametrize("block", [{"type": "duplicate", "refs": ["p1"]},
    {"type": "paragraph", "refs": ["p1"], "duplicateOf": ["p2"]}, {"type": "paragraph", "refs": []}])
def test_duplicate_shape_requires_source_refs_and_only_duplicates_can_target_other_blocks(block):
    with pytest.raises(ValidationError):
        ImportContent(sections=[{"blocks": [block]}])


def test_legacy_refs_and_explicit_blocks_cannot_conflict_in_one_section():
    with pytest.raises(ValidationError, match="either blocks or refs/listItems"):
        ImportContent(sections=[{"refs": ["p1"], "blocks": [{"type": "paragraph", "refs": ["p2"]}]}])


def test_semantic_block_storage_bounds_are_atomic():
    _sources, parts, paragraphs = paragraphs_fixture([f"Detail {i}\n" for i in range(900)])
    before = draft_fixture()
    with pytest.raises(ImportFailure, match="IMPORT_DRAFT_TOO_LARGE"):
        update_content(before, ImportContent(sections=[{"blocks": [
            {"type": "paragraph", "refs": [f"p{i+1}"]} for i in range(900)]}] * 6), paragraphs, parts)
    assert before.sequence == 0 and not before.blocks


@pytest.mark.asyncio
async def test_model_duplicate_interpretation_is_kept_without_a_text_equality_veto(tmp_path, monkeypatch):
    service, identifier, _runtime, _sessions = await selected_service(tmp_path, monkeypatch)
    sources, parts, paragraphs = paragraphs_fixture(["Candidate", "Phone: +7 (123) 456", "Telephone +7 123456"])
    workspace = service.workspaces[identifier]
    workspace.pre_draft.sources = sources
    workspace.pre_draft.selected_source_id = sources[0].id
    saved = update_content(draft_fixture(), ImportContent(sections=[{"category": "identity", "refs": ["p1"]},
        {"category": "contacts", "blocks": [{"type": "paragraph", "refs": ["p2"]},
            {"type": "duplicate", "refs": ["p3"], "duplicateOf": ["p2"]}]}]), paragraphs, parts)
    saved.finished = True
    project = assemble_agent_import(workspace, saved)
    text = " ".join(authored_text(b) for b in project.document.blocks())
    assert "Phone: +7 (123) 456" in text and "Telephone" not in text
    assert project.sources == sources
    assert project.import_coverage[-1].disposition == "excluded"
    await service.shutdown()


def test_explicit_prose_is_not_split_into_lists_or_joined_to_an_adjacent_item():
    sources, blocks = list_fixture(["• Complete item\n", "-X --option\n", "5. Literal prose\n"])
    blocks[0]["type"] = "bullet_list"
    result = format_import_lists({"main": blocks}, sources, lambda _kind, key: str(key),
        model_blocks={"b1", "b2"})["main"]
    assert len(result) == 3
    assert [b["type"] for b in result] == ["bullet_list", "paragraph", "paragraph"]
    assert "".join(s["text"] for s in result[2]["content"]) == "5. Literal prose\n"


@pytest.mark.asyncio
@pytest.mark.parametrize("block_type", ["heading", "paragraph"])
async def test_explicit_identity_stays_in_header_when_fact_wording_differs(tmp_path, monkeypatch, block_type):
    service, identifier, _runtime, _sessions = await selected_service(tmp_path, monkeypatch)
    sources, parts, paragraphs = paragraphs_fixture(["Candidate", "Engineering lead", "Backend engineer"])
    workspace = service.workspaces[identifier]
    workspace.pre_draft.sources = sources
    workspace.pre_draft.selected_source_id = sources[0].id
    saved = update_content(draft_fixture(), ImportContent(profile={"name": "Candidate",
        "title": "Engineering lead; Backend engineer", "refs": ["p1-p3"]},
        sections=[{"category": "identity", "blocks": [{"type": block_type, "refs": ["p1"]},
            {"type": block_type, "refs": ["p2-p3"]}]}]), paragraphs, parts)
    saved.finished = True
    project = assemble_agent_import(workspace, saved)
    assert [authored_text(b) for b in project.document.regions[0].blocks] == ["Candidate", "Engineering lead Backend engineer"]
    assert [b.props.level for b in project.document.regions[0].blocks] == [1, 2]
    assert project.facts.records[0].title.raw_text == "Engineering lead; Backend engineer"
    assert not any(authored_text(b).strip() for r in project.document.regions[1:] for b in r.blocks)
    await service.shutdown()


@pytest.mark.asyncio
async def test_explicit_blocks_are_not_hidden_by_legacy_name_or_contact_deduplication(tmp_path, monkeypatch):
    service, identifier, _runtime, _sessions = await selected_service(tmp_path, monkeypatch)
    sources, parts, paragraphs = paragraphs_fixture(["Candidate", "Phone: 123", "Phone: 123", "Candidate"])
    workspace = service.workspaces[identifier]
    workspace.pre_draft.sources = sources
    workspace.pre_draft.selected_source_id = sources[0].id
    saved = update_content(draft_fixture(), ImportContent(profile={"name": "Candidate", "refs": ["p1"]},
        sections=[{"category": "identity", "blocks": [{"type": "heading", "refs": ["p1"]}]},
            {"category": "contacts", "blocks": [{"type": "paragraph", "refs": ["p2"]}, {"type": "paragraph", "refs": ["p3"]}]},
            {"category": "other", "blocks": [{"type": "paragraph", "refs": ["p4"]}]}]), paragraphs, parts)
    saved.finished = True
    project = assemble_agent_import(workspace, saved)
    text = [authored_text(b) for b in project.document.blocks()]
    assert text.count("Candidate") == text.count("Phone: 123") == 2
    assert all(row.disposition == "placed" for row in project.import_coverage)
    await service.shutdown()


@pytest.mark.asyncio
async def test_tool_checkpoint_reread_and_portable_transfer_keep_duplicate_coverage_without_inference(tmp_path, monkeypatch):
    service, _identifier, runtime, _sessions = await selected_service(tmp_path, monkeypatch)
    raw = b"Company | 2012 -\n2020\nCompany | 2012 - 2020\n"
    workspace = await service.select(raw, "blocks.txt", True)
    identifier = workspace.project_id
    coordinator = ImportCoordinator(service, identifier, "semantic-blocks", service.rules.load(), threading.Event())
    coordinator.begin()
    session = ImportSession(coordinator)
    reopened = None
    try:
        assert (await session.dispatcher.call("read_import", ImportRead()))["nextOffset"] is None
        reply = await session.dispatcher.call("save_import", ImportContent(facts=[{"text": "Likely work period.", "refs": ["p1-p3"]}],
            sections=[{"category": "experience", "blocks": [{"type": "heading", "refs": ["p1-p2"]},
                {"type": "duplicate", "refs": ["p3"], "duplicateOf": ["p1-p2"]}]}]))
        assert reply["durable"] and reply["unassignedRefs"] == []
        read = await session.dispatcher.call("read_import", ImportRead(section="draft"))
        sections = [item["section"] for item in read["items"] if "section" in item]
        assert sections[0]["blocks"][1]["duplicateOf"] == ["p1", "p2"]
        assert (await session.dispatcher.call("finish_import", ImportContentFinish()))["readyForPublication"]
        reopened = RichProjectService(tmp_path / "reopened", RulesStore(tmp_path / "store"), lambda: None)
        assert reopened.read(identifier).pre_draft.import_draft == coordinator.load()
        published = import_workspace(service.read(identifier), coordinator.load(), service.rules.load())
        service.workspaces[identifier] = published
        service.persist(identifier)
        transfer, originals, *_ = open_rich_archive(service.portable(identifier), "block-transfer")
        assert transfer.project.document == published.project.document
        assert transfer.project.facts == published.project.facts
        assert transfer.project.sources == published.project.sources
        assert list(originals.values()) == [raw]
        assert transfer.project.import_coverage[-1].disposition == "excluded"
        assert runtime.calls == []
    finally:
        await session.close()
        if reopened:
            await reopened.shutdown()
        await service.shutdown()


@pytest.mark.asyncio
@pytest.mark.parametrize("queued", [False, True])
@pytest.mark.parametrize("kind", ["pdf", "docx", "markdown"])
async def test_semantic_blocks_publish_in_direct_and_queued_imports_across_formats(tmp_path, monkeypatch, queued, kind):
    from docx import Document
    from resume_api.agent.runtime import RuntimeResult
    from tests.extract.test_structure import pdf
    from tests.support.importing import importer

    texts = ["Candidate", "Company | 2012 -", "2020", "Go, Redis,", "NATS", "Company | 2012 - 2020"]
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    inputs = {"pdf": pdf([(0, 40, 40 + index * 50, text, False) for index, text in enumerate(texts)]),
        "docx": stream.getvalue(), "markdown": ("\n".join(texts) + "\n").encode()}
    raw = inputs[kind]

    async def parse(session, _value, _cancel):
        offset = 0
        while offset is not None:
            reply = await session.dispatcher.call("read_import", ImportRead(offset=offset))
            offset = reply["nextOffset"]
        refs = {p.text.strip(): p.ref for p in session.coordinator.paragraphs}
        reply = await session.dispatcher.call("save_import", ImportContent(
            profile={"name": "Candidate", "refs": [refs[texts[0]]]},
            facts=[{"kind": "employment", "employer": "Company", "periods": [{"start": "2012", "end": "2020"}],
                "technologies": ["Go", "Redis", "NATS"], "refs": [refs[t] for t in texts[1:5]]}],
            sections=[{"category": "identity", "blocks": [{"type": "heading", "refs": [refs[texts[0]]]}]},
                {"category": "experience", "blocks": [{"type": "heading", "refs": [refs[t] for t in texts[1:3]]},
                    {"type": "list_item", "refs": [refs[t] for t in texts[3:5]]},
                    {"type": "duplicate", "refs": [refs[texts[5]]], "duplicateOf": [refs[t] for t in texts[1:3]]}]}]))
        assert reply["durable"] and not reply["unassignedRefs"]
        read = await session.dispatcher.call("read_import", ImportRead(section="draft"))
        assert any(item.get("section", {}).get("blocks") for item in read["items"])
        assert (await session.dispatcher.call("finish_import", ImportContentFinish()))["readyForPublication"]
        return RuntimeResult(state="completed")

    if queued:
        from tests.recovery.test_background_agent import environment
        from tests.recovery.test_background_fixed_jobs import fixed_request, settled
        jobs, factory, *_ = await environment(tmp_path)
        projects = jobs.projects
        projects.provider_factory = lambda: (object(), "processing-model")
        runtime, _sessions = importer(monkeypatch, parse, admission=jobs.admission)
        jobs.runtime = runtime
    else:
        runtime, _sessions = importer(monkeypatch, parse)
        projects = RichProjectService(tmp_path / "operations", RulesStore(tmp_path / "store"),
            lambda: (object(), "fixture"), import_runtime=runtime)
    try:
        selected = await projects.select(raw, "source." + ("md" if kind == "markdown" else kind), False)
        identifier = selected.project_id
        original = selected.pre_draft.sources[0].model_copy(deep=True)
        if queued:
            reply, _ = await jobs.fixed.submit(identifier, fixed_request(jobs, identifier, "blocks", {"kind": "processing"}))
            await settled(jobs, identifier, reply.job.id, publication="saved")
        else:
            generation = projects.start(identifier)
            await projects.tasks[(identifier, generation)]
            assert projects.generations[identifier].status == "completed"
        saved, _ = projects.storage.read_committed(identifier)
        assert saved.project.sources[0] == original
        text = [authored_text(b) for b in saved.project.document.blocks()]
        assert text.count("Company | 2012 - 2020") == 1 and "Go, Redis, NATS" in text
        assert "2020" not in text and "NATS" not in text[1:]
        assert saved.project.import_coverage[-1].disposition == "excluded"
        transfer, originals, *_ = open_rich_archive(projects.portable(identifier), "format-transfer")
        assert transfer.project.document == saved.project.document and transfer.project.sources[0] == original
        assert list(originals.values()) == [raw]
        assert len(runtime.calls) == 1
    finally:
        if queued:
            await jobs.shutdown()
            factory.verify()
            jobs.supervisor.clock.verify()
        else:
            await projects.shutdown()
