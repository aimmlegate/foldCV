"""A job's technologies show as chips: each bound technology gets the tags node's background, separators stay text."""

from docx import Document
from io import BytesIO
import pytest

from packages.resume_core.models.template_document import FieldOccurrence, PresentedText
from packages.resume_core.models.template_facts import ExtensionCell
from packages.resume_core.models.template_workspace import WorkspaceV7
from packages.resume_core.templates.materialize_nodes import DisplayField, render_nodes
from packages.resume_core.templates.presentation_plan import editor_plan
from tests.templates.employment_fixtures import employment_package
from tests.templates.generic_fixtures import generic_package
from tests.templates.test_captured_docx import NS, emitted, xml
from tests.templates.test_materialize import issue_workspace
from tests.templates.test_panel_preparation import _span, _technology, move, prepared
from tests.templates.test_template_view_adapters import render

NAMES = ['Python', 'PostgreSQL', 'Docker']


def cell(name):
    return ExtensionCell(id='cell-' + name, fieldKey='display', revision=1, rawInput=name, typedValue=name, state='admitted',
        origin='source', userCorrected=False, evidence=[], observationIds=[], admissionReceiptIds=[], manualScopeId=None)


def tags_content(package, view):
    declaration = next(block for block in package.definition.blocks if block.id == view)
    fields = [DisplayField(('technologies', str(index), 'display'), 'record-' + name, 'handle-' + name, cell(name))
        for index, name in enumerate(NAMES)]
    blocks, _displayed = render_nodes(package, declaration, 'instance', fields, {}, set())
    [line] = [block for block in blocks if package.recipe.views[view].nodes[block['presentation'].node_index].type == 'tags']
    return [(type(part).__name__, part.text) for part in line['content']]


def test_the_bundled_template_separates_chips_with_a_space_and_an_older_tags_node_keeps_its_comma():
    assert tags_content(generic_package(), 'projects') == [('FieldOccurrence', 'Python'), ('PresentedText', ' '),
        ('FieldOccurrence', 'PostgreSQL'), ('PresentedText', ' '), ('FieldOccurrence', 'Docker')]
    assert [text for _kind, text in tags_content(employment_package(), 'employment-view')] == [
        'Python', ', ', 'PostgreSQL', ', ', 'Docker']
    node = next(node for node in generic_package().recipe.views['projects'].nodes if node.type == 'tags')
    assert (node.style, node.background, node.separator) == ('tag', 'band', ' ')


def stack(separator=', '):
    """The employment fixture's job with three technologies on its tags line; the package declares no separator."""
    before, sources = issue_workspace()
    project = before.project
    job = next(record for record in project.facts.records if record.kind == 'employment')
    python = next(record for record in project.facts.records if record.kind == 'technology')
    added = [_technology('postgres', 'PostgreSQL', job.id), _technology('docker', 'Docker', job.id)]
    project.facts.records.extend(added)
    line = next(block for block in project.document.blocks() if any(isinstance(part, FieldOccurrence)
        and part.record_id == python.id for part in block.content))
    ref = line.presentation
    content = []
    for record in [python, *added]:
        if content:
            content.append(PresentedText(type='text', text=separator, styles={}, presentation=ref))
        content.append(_span(record, ref, '-job'))
    line.content = content
    return WorkspaceV7.model_validate(before.model_dump(by_alias=True)), sources, line.id


def chip_runs(content, names):
    """For each run of the tags line: its text, shading fill, border color and expanded spacing."""
    tree = xml(content)
    [paragraph] = [node for node in tree.xpath('//w:body//w:p', namespaces=NS) if emitted(node) == names]
    return [(emitted(run), run.xpath('string(w:rPr/w:shd/@w:fill)', namespaces=NS),
        run.xpath('string(w:rPr/w:bdr/@w:color)', namespaces=NS), run.xpath('string(w:rPr/w:spacing/@w:val)', namespaces=NS))
        for run in paragraph.xpath('w:r', namespaces=NS)]


def test_the_plan_and_docx_chip_only_the_bound_technologies(tmp_path):
    workspace, _sources, line_id = stack()
    line = next(block for block in workspace.project.document.blocks() if block.id == line_id)
    entry = editor_plan(workspace.project)['blocks'][line_id]
    assert entry['chips'] == {part.binding_id: '#EFEFEF' for part in line.content if isinstance(part, FieldOccurrence)}
    assert len(entry['chips']) == 3
    # The comma separators of an older package stay plain and unshaded; only whitespace between chips is widened.
    assert chip_runs(render(workspace, tmp_path), 'Python, PostgreSQL, Docker') == [
        ('Python', 'EFEFEF', 'EFEFEF', ''), (', ', '', '', ''), ('PostgreSQL', 'EFEFEF', 'EFEFEF', ''),
        (', ', '', '', ''), ('Docker', 'EFEFEF', 'EFEFEF', '')]
    spaced, _sources, _line = stack(' ')
    assert [value[3] for value in chip_runs(render(spaced, tmp_path), 'Python PostgreSQL Docker')] == ['', '80', '', '80', '']


def test_a_user_background_mark_keeps_its_color_on_the_chip(tmp_path):
    workspace, _sources, line_id = stack()
    line = next(block for block in workspace.project.document.blocks() if block.id == line_id)
    line.content[0].styles.background_color = '#ABCDEF'
    workspace = WorkspaceV7.model_validate(workspace.model_dump(by_alias=True))
    assert editor_plan(workspace.project)['blocks'][line_id]['chips'][line.content[0].binding_id] == '#EFEFEF'
    assert chip_runs(render(workspace, tmp_path), 'Python, PostgreSQL, Docker')[0] == ('Python', 'ABCDEF', 'ABCDEF', '')


def test_reordering_technologies_keeps_every_chip_and_the_authored_separators(tmp_path):
    workspace, _sources, line_id = stack()
    moved = prepared(workspace, move(workspace, 'docker', next(record.id for record in workspace.project.facts.records
        if record.kind == 'technology' and record.display.raw_text == 'Python'))).preview.workspace
    line = next(block for block in moved.project.document.blocks() if block.id == line_id)
    assert [part.text for part in line.content] == ['Docker', ', ', 'Python', ', ', 'PostgreSQL']
    chips = editor_plan(moved.project)['blocks'][line_id]['chips']
    assert set(chips) == {part.binding_id for part in line.content if isinstance(part, FieldOccurrence)}
    assert [value[0] for value in chip_runs(render(moved, tmp_path), 'Docker, Python, PostgreSQL') if value[1]] == [
        'Docker', 'Python', 'PostgreSQL']


@pytest.mark.slow
@pytest.mark.asyncio
async def test_a_scripted_generic_import_shows_job_technologies_as_chips_in_the_editor_and_docx(tmp_path):
    from tests.templates.generic_fixtures import generic_import
    _fixture, _result, workspace, _originals = await generic_import(tmp_path / 'import', 'F10')
    blocks = list(workspace.project.document.blocks())
    [line] = [block for block in blocks if [getattr(part, 'text', None) for part in block.content] == ['Python', ' ', 'PostgreSQL', ' ', 'Docker']]
    assert line.presentation.style_token == 'tag'
    plan = editor_plan(workspace.project)
    assert set(plan['blocks'][line.id]['chips']) == {part.binding_id for part in line.content if isinstance(part, FieldOccurrence)}
    assert plan['styles']['tag']['size'] == 9 and plan['styles']['tag']['lineHeight'] == 16
    # The bottom skill groups are join_items, not tags: they stay comma-separated text.
    assert sum('chips' in entry for entry in plan['blocks'].values()) == 1
    content = render(workspace, tmp_path)
    assert chip_runs(content, 'Python PostgreSQL Docker') == [('Python', 'EFEFEF', 'EFEFEF', ''), (' ', '', '', '80'),
        ('PostgreSQL', 'EFEFEF', 'EFEFEF', ''), (' ', '', '', '80'), ('Docker', 'EFEFEF', 'EFEFEF', '')]
    reopened = Document(BytesIO(content))
    assert any(paragraph.text == 'Python PostgreSQL Docker' for table in reopened.tables
        for row in table.rows for table_cell in row.cells for paragraph in table_cell.paragraphs)
