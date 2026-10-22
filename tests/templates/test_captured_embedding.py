"""Actual DOCX parts retain exact captured faces and independently pinned key order."""

from dataclasses import replace
from io import BytesIO
import uuid
from zipfile import ZipFile,ZIP_DEFLATED

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import pytest

from packages.resume_core.rendering.captured_embedding import embed_captured_fonts,obfuscate_font
from packages.resume_core.rendering.captured_fonts import CapturedFontResolver,FontRequest
from packages.resume_core.rendering.template import RenderFailure
from packages.resume_core.templates.compiler import sha256
from tests.templates.test_captured_assets import package


W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P='http://schemas.openxmlformats.org/package/2006/relationships'
T='http://schemas.openxmlformats.org/package/2006/content-types'


@pytest.fixture(scope='module')
def selections():
    captured=package();resolver=CapturedFontResolver(captured)
    return [resolver.select(FontRequest(token)) for token in captured.recipe.style.fonts]


def specimen(selections):
    document=Document()
    for selection in selections:
        run=document.add_paragraph().add_run('Exact captured face: '+selection.font.full_name)
        run.font.name=selection.font.family;run.bold=selection.font.bold;run.italic=selection.font.italic
    output=BytesIO();document.save(output)
    return output.getvalue()


def unpack(raw):
    with ZipFile(BytesIO(raw)) as archive:return {name:archive.read(name) for name in archive.namelist()}


def repack(parts):
    output=BytesIO()
    with ZipFile(output,'w',compression=ZIP_DEFLATED) as archive:
        for name,content in parts.items():archive.writestr(name,content)
    return output.getvalue()


def test_obfuscation_matches_an_independent_known_byte_vector():
    key=uuid.UUID('00112233-4455-6677-8899-aabbccddeeff')
    expected=bytes.fromhex('ffefdfcfbfaf9f8f7f6f5f4f3f2f1f0f' 'efffcfdfafbf8f9f6f7f4f5f2f3f0f1f')
    assert obfuscate_font(bytes(range(32))+b'unchanged tail',key)==expected+b'unchanged tail'


def test_every_actual_embedded_face_relationship_recovers_exact_captured_bytes(selections):
    raw=embed_captured_fonts(specimen(selections),selections);parts=unpack(raw)
    fonts=etree.fromstring(parts['word/fontTable.xml']);rels=etree.fromstring(parts['word/_rels/fontTable.xml.rels'])
    relationships={node.get('Id'):node for node in rels};found=set();variants=set()
    for selection in selections:
        font=selection.font
        family=fonts.xpath('w:font[@w:name=$family]',namespaces={'w':W},family=font.family)
        assert len(family)==1
        tag={(False,False):'embedRegular',(True,False):'embedBold',
            (False,True):'embedItalic',(True,True):'embedBoldItalic'}[(font.bold,font.italic)]
        embedded=family[0].find(qn('w:'+tag));assert embedded is not None
        relationship=relationships[embedded.get(qn('r:id'))]
        assert relationship.get('Type')==R+'/font' and relationship.get('Target')=='fonts/'+font.sha256+'.odttf'
        stored=parts['word/'+relationship.get('Target')]
        key=bytes.fromhex(embedded.get(qn('w:fontKey')).strip('{}').replace('-',''))[::-1]
        original=bytes(stored[index]^key[index%16] for index in range(32))+stored[32:]
        assert original==font.content and sha256(original)==font.sha256
        assert embedded.get(qn('w:subsetted'))=='0'
        found.add(font.sha256);variants.add(tag)
    assert len(found)==10 and len(relationships)==10
    assert variants=={'embedRegular','embedBold','embedItalic','embedBoldItalic'}
    default=etree.fromstring(parts['[Content_Types].xml']).find('{'+T+'}Default[@Extension="odttf"]')
    assert default.get('ContentType')=='application/vnd.openxmlformats-officedocument.obfuscatedFont'
    settings=etree.fromstring(parts['word/settings.xml'])
    assert settings.find(qn('w:embedTrueTypeFonts')).get(qn('w:val'))=='1'
    assert settings.find(qn('w:saveSubsetFonts')).get(qn('w:val'))=='0'
    # Opening with python-docx checks the new package relationships as well as XML parsing.
    reopened=Document(BytesIO(raw));assert len(reopened.paragraphs)==10


def test_face_children_stay_in_schema_order_and_identical_aliases_share_parts(selections):
    regular=next(value for value in selections if value.font.weight==400 and not value.italic)
    alias=replace(regular,font=replace(regular.font,asset_id='same-byte-alias'))
    parts=unpack(embed_captured_fonts(specimen(selections),[*reversed(selections),alias,regular]))
    fonts=etree.fromstring(parts['word/fontTable.xml'])
    family=fonts.xpath('w:font[@w:name="Source Sans Pro"]',namespaces={'w':W})[0]
    actual=[etree.QName(child).localname for child in family if etree.QName(child).localname.startswith('embed')]
    assert actual==['embedRegular','embedBold','embedItalic','embedBoldItalic']
    assert len([name for name in parts if name.endswith('.odttf')])==10


def test_existing_parts_and_relationship_ids_are_preserved(selections):
    original=unpack(specimen(selections));original['word/fonts/retained.bin']=b'preserve exact existing bytes'
    original['word/_rels/fontTable.xml.rels']=(f'<Relationships xmlns="{P}"><Relationship Id="rIdFoldCVCapturedFont1" '
        f'Type="{R}/font" Target="fonts/retained.bin"/></Relationships>').encode()
    result=unpack(embed_captured_fonts(repack(original),selections))
    for name,content in original.items():
        if name not in {'word/fontTable.xml','word/settings.xml','[Content_Types].xml','word/_rels/fontTable.xml.rels'}:
            assert result[name]==content
    rels=etree.fromstring(result['word/_rels/fontTable.xml.rels'])
    assert rels[0].get('Target')=='fonts/retained.bin'
    assert len({node.get('Id') for node in rels})==11
    settings=etree.fromstring(result['word/settings.xml']);tags=[etree.QName(node).localname for node in settings]
    assert tags.index('zoom')<tags.index('embedTrueTypeFonts')<tags.index('saveSubsetFonts')<tags.index('defaultTabStop')


@pytest.mark.parametrize('damage',['hash','rights','same_face','already_embedded','target','content_type'])
def test_embedding_rejects_selected_font_or_existing_package_conflicts(selections,damage):
    raw=specimen(selections);selected=selections;first=selections[0]
    if damage=='hash':selected=[replace(first,font=replace(first.font,content=b'x'*64))]
    elif damage=='rights':selected=[replace(first,font=replace(first.font,embedding_rights=0x0004))]
    elif damage=='same_face':
        content=first.font.content+b'different';other=replace(first.font,content=content,sha256=sha256(content))
        selected=[first,replace(first,font=other)]
    elif damage=='already_embedded':raw=embed_captured_fonts(raw,[first])
    elif damage=='target':
        parts=unpack(raw);parts['word/fonts/'+first.font.sha256+'.odttf']=b'conflict';raw=repack(parts)
    else:
        parts=unpack(raw);types=etree.fromstring(parts['[Content_Types].xml'])
        etree.SubElement(types,'{'+T+'}Default',Extension='odttf',ContentType='conflicting/type')
        parts['[Content_Types].xml']=etree.tostring(types);raw=repack(parts)
    with pytest.raises(RenderFailure,match='FONT_MISMATCH'):embed_captured_fonts(raw,selected)


@pytest.mark.parametrize('damage',['bad_zip','missing_part','duplicate','unsafe','wrong_root','dtd'])
def test_embedding_rejects_malformed_docx_without_returning_partial_artifact(selections,damage):
    parts=unpack(specimen(selections))
    if damage=='bad_zip':raw=b'not a zip'
    elif damage=='duplicate':
        output=BytesIO()
        with pytest.warns(UserWarning,match='Duplicate name'):
            with ZipFile(output,'w') as archive:
                archive.writestr('word/fontTable.xml',b'first');archive.writestr('word/fontTable.xml',b'second')
        raw=output.getvalue()
    else:
        if damage=='missing_part':del parts['word/fontTable.xml']
        elif damage=='unsafe':parts['../escape']=b'invalid'
        elif damage=='wrong_root':parts['word/fontTable.xml']=b'<wrong/>'
        else:parts['word/fontTable.xml']=b'<!DOCTYPE fonts [<!ENTITY content "bad">]><fonts/>'
        raw=repack(parts)
    with pytest.raises(RenderFailure,match='RICH_RENDER_FAILED'):embed_captured_fonts(raw,selections)


def test_embedding_checks_compressed_expanded_and_font_byte_bounds(selections,monkeypatch):
    raw=specimen(selections);parts=unpack(raw)
    monkeypatch.setattr('packages.resume_core.rendering.captured_embedding.DOCX_BYTES',len(raw)-1)
    with pytest.raises(RenderFailure,match='RENDER_OUTPUT_LIMIT'):embed_captured_fonts(raw,selections)
    expanded=sum(map(len,parts.values()))
    monkeypatch.setattr('packages.resume_core.rendering.captured_embedding.DOCX_BYTES',expanded-1)
    with pytest.raises(RenderFailure,match='RENDER_OUTPUT_LIMIT'):embed_captured_fonts(raw,selections)
    monkeypatch.setattr('packages.resume_core.rendering.captured_embedding.DOCX_BYTES',expanded+sum(len(value.font.content) for value in selections)-1)
    with pytest.raises(RenderFailure,match='RENDER_OUTPUT_LIMIT'):embed_captured_fonts(raw,selections)


def test_final_archive_overhead_cannot_exceed_the_output_bound(selections,monkeypatch):
    raw=specimen([]);expanded=sum(map(len,unpack(raw).values()))
    monkeypatch.setattr('packages.resume_core.rendering.captured_embedding.DOCX_BYTES',expanded+1000)
    monkeypatch.setattr('packages.resume_core.rendering.captured_embedding.ZIP_DEFLATED',0)
    with pytest.raises(RenderFailure,match='RENDER_OUTPUT_LIMIT'):embed_captured_fonts(raw,[])
