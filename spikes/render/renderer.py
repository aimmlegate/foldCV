"""Editable DOCX reference layout and LibreOffice PDF conversion spike."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "skills/format-resume/assets"
SOFFICE = ROOT / ".tools/LibreOffice.app/Contents/MacOS/soffice"
APP_FONTS = ROOT / ".tools/LibreOffice.app/Contents/Resources/fonts/truetype"
FONT_FILES = tuple(sorted((ASSETS / "fonts").glob("SourceSansPro-*.ttf")))
FONT_FAMILY = {
    "regular": "Source Sans Pro",
    "light": "Source Sans Pro",
    "semibold": "Source Sans Pro",
    "bold": "Source Sans Pro",
    "black": "Source Sans Pro",
}


def require_environment() -> None:
    if not SOFFICE.is_file():
        raise RuntimeError(f"LibreOffice converter missing: {SOFFICE}")
    if len(FONT_FILES) != 5:
        raise RuntimeError("Five Source Sans Pro font assets are required")
    for source in FONT_FILES:
        installed = APP_FONTS / source.name
        if not installed.is_file() or _sha256(source) != _sha256(installed):
            raise RuntimeError(f"App-local font missing or changed: {installed}; run setup-fonts")


def setup_fonts() -> None:
    """Copy licensed fonts only into the ignored local LibreOffice bundle."""
    if not SOFFICE.is_file():
        raise RuntimeError(f"LibreOffice converter missing: {SOFFICE}")
    if len(FONT_FILES) != 5:
        raise RuntimeError("Five Source Sans Pro font assets are required")
    APP_FONTS.mkdir(parents=True, exist_ok=True)
    for source in FONT_FILES:
        shutil.copy2(source, APP_FONTS / source.name)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _font(run, weight: str = "light", size: float = 10, color: str = "171717") -> None:
    run.font.name = FONT_FAMILY[weight]
    run.font.size = Pt(size)
    run.font.bold = weight in {"semibold", "bold", "black"}
    run.font.color.rgb = RGBColor.from_string(color)
    run.font._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FAMILY[weight])


def _para(container, text: str = "", *, weight: str = "light", size: float = 10,
          before: float = 0, after: float = 3, keep_next: bool = False,
          bullet: bool = False):
    p = container.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_with_next = keep_next
    p.paragraph_format.widow_control = True
    if bullet:
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        _font(p.add_run("•  "), "regular", size)
    _font(p.add_run(text), weight, size)
    return p


def _fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _cell_margin(cell, top: int = 0, start: int = 0, bottom: int = 0, end: int = 0) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = qn(f"w:{side}")
        element = margins.find(tag)
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _section(cell, title: str, *, bar: bool = False) -> None:
    if bar:
        table = cell.add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = Inches(5.78)
        panel = table.cell(0, 0)
        _fill(panel, "EEEEEE")
        _cell_margin(panel, top=40, start=80, bottom=40, end=40)
        p = panel.paragraphs[0]
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(0)
        _font(p.add_run(title), "bold", 14)
        trailing = cell.paragraphs[-1]
        if not trailing.text:
            trailing._element.getparent().remove(trailing._element)
    else:
        _para(cell, title, weight="bold", size=14, after=8, keep_next=True)


def _remove_empty_first(cell) -> None:
    p = cell.paragraphs[0]
    if not p.text:
        p._element.getparent().remove(p._element)


def _header(doc, candidate: dict) -> None:
    header = doc.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Inches(6.27)
    header.columns[1].width = Inches(1.5)
    for cell in header.rows[0].cells:
        _cell_margin(cell)
    p = header.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    _font(p.add_run(candidate["name"]), "black", 18)
    _para(header.cell(0, 0), candidate["title"], weight="semibold", size=16, after=0)
    p = header.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(ASSETS / "images/foldcv.png"), width=Inches(1.42))
    _para(doc, after=2)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:color"), "888888")
    borders.append(bottom)
    p_pr.append(borders)


def _sidebar(cell, data: dict) -> None:
    _remove_empty_first(cell)
    if data.get("focus"):
        _section(cell, "ФОКУС")
        for item in data["focus"]:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = bool(item.get("duration"))
            icon = item.get("icon")
            if icon:
                icon_path = ASSETS / "images" / f"{icon}.png"
                if not icon_path.is_file():
                    icon_path = ASSETS / "devicon" / f"{icon}.png"
                if not icon_path.is_file():
                    raise ValueError(f"Unknown focus icon: {icon}")
                p.add_run().add_picture(
                    str(icon_path), width=Inches(0.3), height=Inches(0.3))
            _font(p.add_run("  " + item["label"]), "semibold", 10)
            if item.get("duration"):
                duration = _para(cell, item["duration"], size=9, after=7)
                duration.paragraph_format.left_indent = Inches(0.60)
    for section in data.get("sidebar", []):
        if section.get("items"):
            _section(cell, section["title"])
            for item in section["items"]:
                _para(cell, item, after=7)


def _main(cell, data: dict) -> None:
    _remove_empty_first(cell)
    if data.get("summary"):
        _section(cell, "О СЕБЕ")
        _para(cell, data["summary"], after=11)
    if data.get("projects_title"):
        _section(cell, "ПРОЕКТЫ", bar=True)
    for project in data.get("projects", []):
        _para(cell, project["company"], weight="semibold", size=12, before=9, after=1, keep_next=True)
        _para(cell, project["period"], weight="regular", after=1, keep_next=True)
        if project.get("stack"):
            _para(cell, project["stack"], weight="regular", after=1, keep_next=True)
        if project.get("description"):
            _para(cell, project["description"], after=4, keep_next=True)
        _para(cell, "Роль: " + project["role"], after=4, keep_next=bool(project.get("tasks") or project.get("results")))
        for title, key in (("Задачи:", "tasks"), ("Личные результаты:", "results")):
            if project.get(key):
                _para(cell, title, after=1, keep_next=True)
                for item in project[key]:
                    _para(cell, item, bullet=True, after=2)
    if data.get("achievements"):
        _para(cell, "Ключевые достижения:", weight="semibold", before=9, keep_next=True)
        for item in data["achievements"]:
            _para(cell, item, bullet=True)
    for section in data.get("skill_groups", []):
        _section(cell, section["title"], bar=True)
        _para(cell, section["items"], after=8)


def _lines(text: str, width: float, font: str = "regular", size: float = 10) -> int:
    count = 0
    for explicit_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        count += 1
        current = ""
        for word in explicit_line.split():
            if pdfmetrics.stringWidth(word, font, size) > width:
                raise ValueError("Unbreakable text exceeds the column width")
            trial = f"{current} {word}".strip()
            if current and pdfmetrics.stringWidth(trial, font, size) > width:
                count += 1
                current = word
            else:
                current = trial
    return count


def _measure(text: str, width: float, *, font: str = "regular", size: float = 10,
             before: float = 0, after: float = 3) -> float:
    return _lines(text, width, font, size) * size * 1.15 + before + after


def _page_budget(index: int) -> float:
    return 615 if index == 0 else 705


def _append_with_budget(pages: list[list], item, height: float) -> None:
    if height > _page_budget(1):
        raise ValueError("One resume block exceeds the page budget")
    page = pages[-1]
    used = sum(existing[1] for existing in page)
    if page and used + height > _page_budget(len(pages) - 1):
        pages.append([])
    if height > _page_budget(len(pages) - 1):
        raise ValueError("Resume block cannot fit on its planned page")
    pages[-1].append((item, height))


def _paginate_main(data: dict) -> list[dict]:
    pages: list[list] = [[]]
    summary = data["candidate"]["summary"]
    _append_with_budget(pages, ("summary", summary), 35 + _measure(summary, 410, after=11))
    for index, project in enumerate(data.get("projects", [])):
        height = 17 if index == 0 else 0
        height += _measure(project["company"], 410, font="bold", size=12, before=9, after=1)
        height += _measure(project["period"], 410, after=1)
        for field in ("stack", "description"):
            if project.get(field):
                height += _measure(project[field], 410, after=4)
        height += _measure("Роль: " + project["role"], 410, after=4)
        for title, key in (("Задачи:", "tasks"), ("Личные результаты:", "results")):
            if project.get(key):
                height += _measure(title, 410, after=1)
                height += sum(_measure("•  " + line, 385, after=2) for line in project[key])
        _append_with_budget(pages, ("project", (index, project)), height)
    if data.get("achievements"):
        height = 20 + sum(_measure("•  " + line, 385) for line in data["achievements"])
        _append_with_budget(pages, ("achievements", data["achievements"]), height)
    for group in data.get("skill_groups", []):
        if not group["items"].strip():
            continue
        height = 34 + _measure(group["items"], 410, after=8)
        _append_with_budget(pages, ("skill_group", group), height)
    result = []
    for page in pages:
        content = {"projects": [], "achievements": [], "skill_groups": []}
        for (kind, value), _ in page:
            if kind == "summary":
                content["summary"] = value
            elif kind == "project":
                index, project = value
                if index == 0:
                    content["projects_title"] = True
                content["projects"].append(project)
            elif kind == "achievements":
                content["achievements"] = value
            else:
                content["skill_groups"].append(value)
        result.append(content)
    return result


def _paginate_side(data: dict) -> list[dict]:
    pages: list[list] = [[]]
    for index, item in enumerate(data.get("focus", [])):
        height = 31 if index == 0 else 0
        height += max(32, _measure(item["label"], 98, font="bold", after=0)
                      + (_measure(item.get("duration", ""), 98, size=9) if item.get("duration") else 0)) + 7
        _append_with_budget(pages, ("focus", item), height)
    for section in data.get("sidebar", []):
        for item in section.get("items", []):
            height = _measure(item, 105, after=7) + 5
            if not any(entry[0][0] == "sidebar" and entry[0][1][0] == section["title"] for entry in pages[-1]):
                height += 34
            _append_with_budget(pages, ("sidebar", (section["title"], item)), height)
    result = []
    for page in pages:
        content = {"focus": [], "sidebar": []}
        for (kind, value), _ in page:
            if kind == "focus":
                content["focus"].append(value)
            else:
                title, item = value
                if not content["sidebar"] or content["sidebar"][-1]["title"] != title:
                    content["sidebar"].append({"title": title, "items": []})
                content["sidebar"][-1]["items"].append(item)
        result.append(content)
    return result


def render_docx(data: dict, target: Path) -> None:
    require_environment()
    pdfmetrics.registerFont(TTFont("regular", str(ASSETS / "fonts/SourceSansPro-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("bold", str(ASSETS / "fonts/SourceSansPro-Bold.ttf")))
    if _lines(data["candidate"]["name"], 450, "bold", 18) > 1:
        raise ValueError("Candidate display name exceeds one header line")
    if _lines(data["candidate"]["title"], 450, "bold", 16) > 2:
        raise ValueError("Candidate title exceeds two header lines")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.2677)
    sec.page_height = Inches(11.6929)
    sec.top_margin = Inches(0.23)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.3)
    sec.right_margin = Inches(0.3)
    sec.header_distance = Inches(0.2)
    sec.footer_distance = Inches(0.25)
    normal = doc.styles["Normal"]
    normal.font.name = "Source Sans Pro"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    _header(doc, data["candidate"])
    side_pages, main_pages = _paginate_side(data), _paginate_main(data)
    for index in range(max(len(side_pages), len(main_pages))):
        if index:
            doc.add_page_break()
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        row_pr = table.rows[0]._tr.get_or_add_trPr()
        split = OxmlElement("w:cantSplit")
        split.set(qn("w:val"), "0")
        row_pr.append(split)
        table.columns[0].width = Inches(1.85)
        table.columns[1].width = Inches(5.78)
        side, main = table.rows[0].cells
        side.width = Inches(1.85)
        main.width = Inches(5.78)
        _cell_margin(side, start=100, end=160)
        _cell_margin(main, start=40, end=0)
        side.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        main.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if index < len(side_pages):
            _sidebar(side, side_pages[index])
        if index < len(main_pages):
            _main(main, main_pages[index])
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(footer.add_run(data.get("footer", "mail@example.com")), "light", 12)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def convert_pdf(docx_path: Path, pdf_path: Path, *, timeout: int = 30) -> None:
    require_environment()
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="foldcv-lo-") as temp:
        temp_path = Path(temp)
        profile = (temp_path / "profile").as_uri()
        out = temp_path / "output"
        out.mkdir()
        command = [str(SOFFICE), f"-env:UserInstallation={profile}", "--headless",
                   "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(out), str(docx_path)]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=timeout,
                                    env={**os.environ, "SAL_USE_VCLPLUGIN": "osx"}, check=False)
        except subprocess.TimeoutExpired as error:
            raise RuntimeError(f"LibreOffice conversion timed out after {timeout}s") from error
        converted = out / (docx_path.stem + ".pdf")
        if result.returncode or not converted.is_file():
            raise RuntimeError(f"LibreOffice conversion failed ({result.returncode}): {result.stderr[-500:]}")
        shutil.copy2(converted, pdf_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("setup-fonts", "render"))
    parser.add_argument("fixture", nargs="?", type=Path)
    parser.add_argument("output", nargs="?", type=Path)
    args = parser.parse_args()
    if args.command == "setup-fonts":
        setup_fonts()
        print(f"Installed {len(FONT_FILES)} fonts in {APP_FONTS}")
        return
    if not args.fixture or not args.output:
        parser.error("render requires fixture JSON and output directory")
    data = json.loads(args.fixture.read_text(encoding="utf-8"))
    docx_path = args.output / f"{args.fixture.stem}.docx"
    pdf_path = args.output / f"{args.fixture.stem}.pdf"
    render_docx(data, docx_path)
    convert_pdf(docx_path, pdf_path)
    print(f"{docx_path}\n{pdf_path}")


if __name__ == "__main__":
    main()
