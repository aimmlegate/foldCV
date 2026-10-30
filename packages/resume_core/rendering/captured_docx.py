"""Editable DOCX from exact authored content and captured template assets."""

from dataclasses import replace
import hashlib
from importlib.metadata import version
from io import BytesIO
from pathlib import Path
import platform
import re

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.shared import Twips

from ..models.template_document import PresentedText, TemplateParagraphBlock
from ..models.template_render import TemplateRendererIdentity, TemplateRenderRequestV1
from ..rich_limits import DOCX_BYTES
from ..templates.compiler import ADDED_CAPABILITIES, CAPABILITIES, canonical
from ..templates.focus import icon_catalog, row_icon
from ..templates.presentation_plan import chip_background, focus_row, hyperlink_target
from ..templates.materialize_nodes import reference
from .captured_assets import CapturedAssets, decode_logo, decode_png
from .captured_embedding import embed_captured_fonts
from .captured_layout import BodyGeometry, CapturedMetrics, half_points, native_units
from .captured_styles import CapturedStyleResolver, authored_spans, resize_run, visible_text
from .rich_docx import Writer, element, fill
from .template import RenderFailure

ROOT = Path(__file__).resolve().parents[3]
RENDERER_VERSION = '1.0.0'
# Word draws no rounded run background: a chip is shaded text with a border of its own color as padding, and a
# whitespace separator between two chips is widened so the borders keep a visible gap.
CHIP_BORDER = 12
CHIP_GAP = 4
BLOCKS = {'section', 'paragraph', 'heading', 'bullet_list', 'numbered_list'}
NODES = {'text', 'join_fields', 'join_items', 'period', 'tags', 'bullets', 'link', 'focus_entry'}


def renderer_identity():
    sources = sorted((ROOT / 'packages/resume_core').rglob('*.py')) + [ROOT / 'requirements.txt']
    manifest = {'sources': {path.relative_to(ROOT).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sources}, 'runtime': {'python': platform.python_version(), **{name: version(name)
        for name in ('fonttools', 'python-docx', 'lxml', 'pydantic', 'jsonschema', 'PyYAML', 'PyMuPDF', 'Pillow')}}}
    return TemplateRendererIdentity(id='captured-docx', version=RENDERER_VERSION,
        sha256=hashlib.sha256(canonical(manifest)).hexdigest())


def resolved_node(resolver, ref):
    resolver._reference(ref)
    node = resolver.views[ref.view_key].nodes[ref.node_index] if ref.node_index is not None else None
    if node is not None and node.type not in NODES:
        raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
    return node


def native_text(run, text):
    for part in re.split(r'(\r\n|\r|\n|\t)', text):
        if part in {'\r\n', '\r', '\n'}:
            run._r.append(element('w:br'))
        elif part == '\t':
            run._r.append(element('w:tab'))
        elif part:
            node = element('w:t')
            node.set(qn('xml:space'), 'preserve')
            node.text = part
            run._r.append(node)


def run_properties(style):
    font = style.selection.font
    props = element('w:rPr')
    props.append(element('w:rFonts', ascii=font.family, hAnsi=font.family, eastAsia=font.family, cs=font.family))
    for tag, value in [('b', font.bold), ('bCs', font.bold), ('i', font.italic), ('iCs', font.italic),
            ('strike', style.strike)]:
        props.append(element('w:' + tag, val='1' if value else '0'))
    props.append(element('w:color', val=style.color.removeprefix('#')))
    if style.gap:
        props.append(element('w:spacing', val=native_units(style.gap)))
    size = half_points(style.size)
    props.append(element('w:sz', val=size))
    props.append(element('w:szCs', val=size))
    props.append(element('w:u', val='single' if style.underline else 'none'))
    if style.chip:
        props.append(element('w:bdr', val='single', sz=CHIP_BORDER, space=0, color=style.background.removeprefix('#')))
    if style.background is not None:
        fill(props, style.background)
    return props


class CapturedWriter:
    compact = Writer.compact

    def floating_image(self, paragraph, path, width, height, x, y, *, label, horizontal='page', vertical='page', right=None):
        x, y, width, height = [native_units(value, 12700, positive=index > 1)
            for index, value in enumerate((x, y, width, height))]
        if x + width > (right or self.geometry.width) * 635 or y + height > self.geometry.height * 635:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        Writer.floating_image(self, paragraph, path, width / 12700, height / 12700, x / 12700, y / 12700,
            horizontal, vertical, label=label)
        anchor = paragraph._p.xpath('.//wp:anchor')[-1]
        anchor.find(qn('wp:positionH')).find(qn('wp:posOffset')).text = str(x)
        anchor.find(qn('wp:positionV')).find(qn('wp:posOffset')).text = str(y)
        for extent in [anchor.find(qn('wp:extent')), *anchor.iter(qn('a:ext'))]:
            extent.set('cx', str(width))
            extent.set('cy', str(height))

    def clear_cell(self, cell):
        Writer.clear_cell(self, cell)
        margins = cell._tc.get_or_add_tcPr().find(qn('w:tcMar'))
        margins[:] = [element('w:' + side, w=0, type='dxa')
            for side in ('top', 'start', 'left', 'bottom', 'end', 'right')]

    def __init__(self, request):
        self.request = request
        self.recipe = request.package.recipe
        self.geometry = BodyGeometry.from_recipe(self.recipe)
        if set(self.recipe.capabilities) - set(CAPABILITIES + ADDED_CAPABILITIES):
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        self.resolver = CapturedStyleResolver(request.package)
        self.metrics = CapturedMetrics()
        self.document = Document()
        self.numbering = self.document.part.numbering_part.element
        self.next_numbering = 6000
        self.next_drawing = 100000
        self.selections = {}
        self.widths = {}
        self.regions = {}
        def assign_width(blocks, width, region):
            for block in blocks:
                self.widths[block.id] = native_units(width, positive=True)
                self.regions[block.id] = region
                assign_width(block.children, width, region)
        layout = {region.id: region for region in self.recipe.layout.regions}
        for region in request.document.regions:
            # Without a sidebar column, blocks moved into the sidebar flow at the main width before the main ones.
            assign_width(region.blocks, layout.get(region.role, layout['main']).width, region.role)
        self.runs = {}
        self.links = {}
        self.paragraphs = {}
        self.logo = None
        # The block in work, so that a refusal can name it.
        self.block = None

    def prepare(self):
        request = self.request
        for block in request.document.blocks():
            self.block = block
            if block.type not in BLOCKS or block.type == 'section' and block.props.presentation == 'focus':
                raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
            resolved_node(self.resolver, block.presentation)
            if not isinstance(block.content, list):
                raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
            for span in block.content:
                if span.type not in {'text', 'fact_field', 'manual_scope', 'bound_link', 'link'}:
                    raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
                if span.type in {'bound_link', 'link'} and span.presentation is not None:
                    resolved_node(self.resolver, span.presentation)
            actual, links = [], []
            for span, link in authored_spans(block):
                resolved_node(self.resolver, span.presentation or block.presentation)
                text = visible_text(span.text)
                if not text:
                    continue
                style = self.resolver.run(block, span)
                chip = chip_background(block, span, self.recipe) if link is None else None
                if chip is not None:
                    style = replace(style, background=style.background or chip, chip=True)
                self.remember(style)
                actual.append((text, style))
                links.append(link)
            for index in range(1, len(actual) - 1):
                (_before, left), (text, style), (_after, right) = actual[index - 1:index + 2]
                if left.chip and right.chip and not style.chip and text.isspace() and not set(text) & set('\r\n\t'):
                    actual[index] = (text, replace(style, gap=CHIP_GAP / len(text)))
            self.runs[block.id] = actual
            self.links[block.id] = links
            emitted = (block.type != 'section' or block.content) and (self.regions[block.id] != 'header' or actual)
            self.paragraphs[block.id] = self.resolver.paragraph(block) if emitted else None
            if (self.paragraphs[block.id] is not None and self.regions[block.id] != 'header' and
                    native_units(self.paragraphs[block.id][0], positive=True) > self.geometry.bottom - self.geometry.continuation_top):
                raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
            if not actual and emitted:
                self.remember(self.generated(block, ''))
        self.block = None
        self.rows, self.icons, self.placed = {}, {}, []
        catalog = None
        for block in request.document.blocks():
            row = focus_row(block, self.recipe)
            if not row:
                continue
            if self.regions[block.id] == 'header':
                raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
            try:
                catalog = icon_catalog(request.package) if catalog is None else catalog
            except ValueError:
                raise RenderFailure('ASSET_MISMATCH') from None
            found = row_icon(request.package, request.facts, row, catalog, block.props.icon)
            icon = found[0] if found else None
            if found and icon not in self.icons:
                self.icons[icon] = decode_png(found[1], 'icon')
            self.rows[block.id] = (row, icon)
        self.configure_page()

    def remember(self, style):
        half_points(style.size)
        native_units(style.line_height, positive=True)
        self.selections[style.selection.font.sha256] = style.selection

    def generated(self, block, text):
        span = PresentedText(type='text', text=text, styles={}, presentation=None)
        return self.resolver.run(block, span)

    def configure_page(self):
        page, section = self.geometry, self.document.sections[0]
        section.page_width, section.page_height = Twips(page.width), Twips(page.height)
        section.top_margin, section.bottom_margin = Twips(page.continuation_top), Twips(page.height - page.bottom)
        section.left_margin, section.right_margin = Twips(page.left), Twips(page.right)
        section.header_distance = section.footer_distance = Twips(0)
        section.different_first_page_header_footer = True
        self.document.core_properties.author = ''
        self.document.core_properties.last_modified_by = ''
        normal = self.document.styles['Normal'].paragraph_format
        normal.space_before = normal.space_after = Twips(0)
        normal.keep_together = normal.keep_with_next = normal.widow_control = False

    def emit_run(self, paragraph, text, style):
        self.remember(style)
        run = paragraph.add_run()
        run._r.append(run_properties(style))
        native_text(run, text)
        return run

    def paragraph(self, container, block, *, header=False, depth=0, sequence=None):
        self.block = block
        paragraph = container.add_paragraph()
        height, alignment = self.paragraphs[block.id]
        props = paragraph._p.get_or_add_pPr()
        props.append(element('w:spacing', before=0, after=0, line=native_units(height, positive=True), lineRule='exact'))
        props.append(element('w:jc', val='both' if alignment == 'justify' else alignment))
        props.append(element('w:widowControl', val=0))
        props.append(element('w:keepLines', val=0))
        is_heading = block.type == 'heading'
        props.append(element('w:keepNext', val=int(is_heading and self.recipe.layout.pagination.keep_heading_with_next)))
        view = self.resolver.views[block.presentation.view_key]
        band = is_heading and block.presentation.node_index is None and view.heading == 'band'
        node = resolved_node(self.resolver, block.presentation)
        focus = node is not None and node.type == 'focus_entry'
        inset = 0 if header or band else view.inset + (self.recipe.style.focus.text_x if focus else 0)
        left = native_units(inset)
        if left >= self.widths[block.id]:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        props.append(element('w:ind', left=left, right=0))
        if band:
            band_style = self.recipe.style.section_band
            usable = self.widths[block.id] / 20 - band_style.text_inset
            if (usable <= 0 or height > band_style.height or
                    native_units(band_style.height, positive=True) > self.geometry.bottom - self.geometry.continuation_top or
                    any(any(control in text for control in '\r\n\t') for text, _style in self.runs[block.id]) or
                    self.metrics.paragraph_height(self.runs[block.id], usable, height) > height):
                raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
            props.find(qn('w:ind')).set(qn('w:firstLine'), str(native_units(band_style.text_inset)))
            fill(props, self.recipe.style.colors[self.recipe.style.section_band.background])
            spacing = props.find(qn('w:spacing'))
            spacing.set(qn('w:line'), str(native_units(band_style.height, positive=True)))
        if is_heading:
            props.append(element('w:outlineLvl', val=block.props.level - 1))
        if block.type in {'bullet_list', 'numbered_list'}:
            if header:
                raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
            marker_height = self.list_marker(paragraph, block, inset, depth, sequence)
            line = native_units(max(height, marker_height), positive=True)
            if line > self.geometry.bottom - self.geometry.continuation_top:
                raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
            props.find(qn('w:spacing')).set(qn('w:line'), str(line))
        if not self.runs[block.id]:
            props.append(run_properties(self.generated(block, '')))
        previous_cr, current, hyperlink = False, None, None
        for (text, style), link in zip(self.runs[block.id], self.links[block.id], strict=True):
            run = self.emit_run(paragraph, text[1:] if previous_cr and text.startswith('\n') else text, style)
            previous_cr = text.endswith('\r')
            target = hyperlink_target(link.href) if link is not None else None
            if target is None:
                current = None
                continue
            # Consecutive runs of one link share its element; equal targets share one relationship.
            if link is not current:
                current, hyperlink = link, element('w:hyperlink')
                hyperlink.set(qn('r:id'), paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                paragraph._p.append(hyperlink)
            hyperlink.append(run._r)
        self.block = None
        return paragraph

    def list_marker(self, paragraph, block, inset, depth, sequence):
        if depth > 8:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        bullets = self.recipe.style.bullets
        if bullets.marker_offset > bullets.text_offset:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        numbered = block.type == 'numbered_list'
        start = block.props.start if numbered else None
        index = (start or (sequence.get('index', 0) + 1 if sequence.get('kind') == 'numbered_list' else 1)) if numbered else 1
        if not 1 <= index <= 2147483647:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        marker = f'{index}.' if numbered else bullets.marker
        style = resize_run(self.generated(block, marker), bullets.marker_size)
        self.remember(style)
        left = native_units(inset + bullets.text_offset * (depth + 1))
        marker_left = native_units(inset + bullets.text_offset * depth + bullets.marker_offset)
        if left >= self.widths[block.id]:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        signature = (numbered, style.selection.font.family, style.selection.font.sha256,
            style.size, style.color, left, marker_left)
        identifier = sequence.get('num') if start is None and sequence.get('signature') == signature else None
        if identifier is None:
            identifier = self.next_numbering
            self.next_numbering += 1
            abstract = element('w:abstractNum', abstractNumId=identifier)
            abstract.append(element('w:multiLevelType', val='singleLevel'))
            level = element('w:lvl', ilvl=0)
            level.append(element('w:start', val=index))
            level.append(element('w:numFmt', val='decimal' if numbered else 'bullet'))
            level.append(element('w:lvlText', val='%1.' if numbered else marker))
            level.append(element('w:lvlJc', val='left'))
            props = element('w:pPr')
            tabs = element('w:tabs')
            tabs.append(element('w:tab', val='num', pos=left))
            props.append(tabs)
            props.append(element('w:ind', left=left, hanging=left - marker_left))
            level.append(props)
            level.append(run_properties(style))
            abstract.append(level)
            first_num = self.numbering.find(qn('w:num'))
            if first_num is None:
                self.numbering.append(abstract)
            else:
                first_num.addprevious(abstract)
            number = element('w:num', numId=identifier)
            number.append(element('w:abstractNumId', val=identifier))
            self.numbering.append(number)
        sequence.update(num=identifier, index=index, signature=signature, kind=block.type)
        props = paragraph._p.get_or_add_pPr()
        existing = props.find(qn('w:ind'))
        existing.set(qn('w:left'), str(left))
        existing.set(qn('w:hanging'), str(left - marker_left))
        number = element('w:numPr')
        number.append(element('w:ilvl', val=0))
        number.append(element('w:numId', val=identifier))
        props.insert(0, number)
        return style.line_height

    def blocks(self, container, blocks, depth=0):
        sequence = {}
        for block in blocks:
            if block.type not in {'bullet_list', 'numbered_list'}:
                sequence.clear()
            if block.type != 'section' or block.content:
                self.paragraph(container, block, depth=depth, sequence=sequence)
            row, icon = self.rows.get(block.id, ([], None))
            if row:
                self.focus(container, block, row, icon, depth)
            self.blocks(container, block.children[len(row):], depth + (block.type in {'bullet_list', 'numbered_list'}))
            if block.type == 'section' and block.children:
                gap = self.resolver.views[block.presentation.view_key].item_gap
                if gap:
                    self.spacer(container, native_units(gap, positive=True))

    def focus(self, container, section, row, icon, depth):
        """Focus text centered in at least the minimum row height; the decorative icon centered in its box."""
        style = self.recipe.style.focus
        view = self.resolver.views[row[0].presentation.view_key]
        text = sum(native_units(self.metrics.paragraph_height(self.runs[block.id],
            self.widths[block.id] / 20 - view.inset - style.text_x, self.paragraphs[block.id][0]), positive=True) for block in row)
        height = max(native_units(style.minimum_row_height, positive=True), text)
        before = (height - text) // 2
        anchor = self.spacer(container, before) if before else None
        paragraphs = [self.paragraph(container, block, depth=depth, sequence={}) for block in row]
        # One entry stays on one page: its leading spacer and lines keep with what follows.
        for paragraph in [*([anchor] if anchor is not None else []), *paragraphs[:-1]]:
            paragraph.paragraph_format.keep_with_next = True
        if height - text - before:
            self.spacer(container, height - text - before)
        if icon is not None:
            image = self.icons[icon]
            scale = min(style.icon_width / image.width, style.icon_height / image.height)
            width, image_height = image.width * scale, image.height * scale
            self.placed.append((anchor if anchor is not None else paragraphs[0], icon, width, image_height,
                view.inset + style.icon_x + (style.icon_width - width) / 2, (height / 20 - image_height) / 2, self.widths[section.id]))

    def spacer(self, container, height):
        paragraph = container.add_paragraph()
        self.compact(paragraph, height / 20)
        paragraph.paragraph_format.line_spacing = Twips(height)
        paragraph.paragraph_format.widow_control = False
        return paragraph

    def frame(self, paragraph, x, y, width):
        native_x, native_y = native_units(x), native_units(y)
        native_width = native_units(width, positive=True)
        if native_x + native_width > self.geometry.width or native_y >= self.geometry.height:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        paragraph._p.get_or_add_pPr().insert(0, element('w:framePr', wrap='around',
            hAnchor='page', vAnchor='page', x=native_x, y=native_y,
            w=native_width, hSpace=0, vSpace=0))

    def header(self):
        header = self.document.sections[0].first_page_header
        self.compact(header.paragraphs[0], .05)
        recipe = self.recipe.style.header
        region = self.recipe.layout.regions[0]
        # Source fallback lines stack from the first slot; growth counts from the last slot, as if they filled both.
        first, last = min(recipe.name_top, recipe.profession_top), max(recipe.name_top, recipe.profession_top)
        cursor = 0.0
        growth = 0.0
        for block in self.request.document.regions[0].blocks:
            for child in self.header_blocks(block):
                if not self.runs[child.id]:
                    continue
                refs = [span.presentation or child.presentation for span, _link in authored_spans(child) if visible_text(span.text)]
                if self.header_fallback(child, refs):
                    top, designed = first, last
                else:
                    slots = {getattr(resolved_node(self.resolver, ref), 'header_slot', None) for ref in refs}
                    if len(slots) != 1 or None in slots:
                        raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
                    top = designed = getattr(recipe, next(iter(slots)) + '_top')
                height = self.metrics.paragraph_height(self.runs[child.id], region.width, self.paragraphs[child.id][0])
                paragraph = self.paragraph(header, child, header=True)
                y = max(top, cursor)
                self.frame(paragraph, region.x, y, region.width)
                cursor = y + height
                growth = max(growth, cursor - designed - self.paragraphs[child.id][0])
        # A final ordinary paragraph terminates the last independent text frame.
        self.spacer(header, 1)
        if cursor > self.recipe.style.page.body_bottom or self.recipe.style.page.first_body_top + growth >= self.recipe.style.page.body_bottom:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        return growth

    def header_fallback(self, block, refs):
        """A paragraph of source wording a header block kept: no view node, the block's own fallback style."""
        if block.type != 'paragraph' or not refs or any(resolved_node(self.resolver, ref) is not None for ref in refs):
            return False
        keys = {ref.view_key for ref in refs}
        if len(keys) != 1:
            return False
        view = self.resolver.views[next(iter(keys))]
        style = view.fallback_style or (view.nodes[0].style if view.nodes else None)
        return all(ref.style_token == style for ref in refs)

    def header_blocks(self, block):
        if block.type != 'section' or block.content:
            yield block
        for child in block.children:
            yield from self.header_blocks(child)

    def footer(self):
        recipe = self.recipe.style.footer
        if not recipe.text:
            return
        ref = reference(self.request.package, next(iter(self.recipe.views)), recipe.style, 'captured-footer')
        block = TemplateParagraphBlock(type='paragraph', id='captured-footer', revision=1, props={},
            content=[PresentedText(type='text', text=recipe.text, styles={}, presentation=None)],
            children=[], provenance=[], presentation=ref)
        style = self.generated(block, recipe.text)
        width = min(recipe.center_x, self.recipe.style.page.width - recipe.center_x) * 2
        top = recipe.baseline_y - self.metrics.ascent(style)
        distance = self.recipe.style.page.height - recipe.baseline_y - (style.line_height - self.metrics.ascent(style))
        if top < 0 or width <= 0 or recipe.baseline_y > self.recipe.style.page.height or distance < 0:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        self.document.sections[0].footer_distance = Twips(native_units(distance))
        left = -self.geometry.left + native_units(recipe.center_x - width / 2)
        right = self.geometry.width - self.geometry.right - native_units(recipe.center_x + width / 2)
        for footer in (self.document.sections[0].footer, self.document.sections[0].first_page_footer):
            paragraph = footer.paragraphs[0]
            self.compact(paragraph, style.line_height)
            paragraph.paragraph_format.line_spacing = Twips(native_units(style.line_height, positive=True))
            paragraph._p.get_or_add_pPr().append(element('w:jc', val='center'))
            paragraph._p.get_or_add_pPr().append(element('w:ind', left=left, right=right))
            self.emit_run(paragraph, recipe.text, style)

    def first_page_spacer(self, growth):
        total = self.geometry.first_top - self.geometry.continuation_top + native_units(growth)
        divider = self.recipe.style.header.divider
        if divider is None:
            if total:
                self.spacer(self.document, total)
            return
        page = self.recipe.style.page
        if not page.continuation_body_top <= divider.y <= page.first_body_top:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        x1, x2 = native_units(divider.x1), native_units(divider.x2)
        thickness = native_units(divider.width, 8, positive=True, maximum=96)
        target = native_units(divider.y + growth) - self.geometry.continuation_top
        before = max(0, target - 1)
        after = total - before - 1
        if x1 >= x2 or x2 > self.geometry.width or thickness < 2 or after < 0:
            raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
        if before:
            self.spacer(self.document, before)
        paragraph = self.spacer(self.document, 1)
        props = paragraph._p.get_or_add_pPr()
        props.append(element('w:ind', left=x1 - self.geometry.left,
            right=self.geometry.width - self.geometry.right - x2))
        borders = element('w:pBdr')
        borders.append(element('w:bottom', val='single', color=self.recipe.style.colors[divider.color][1:],
            sz=thickness, space=0))
        props.append(borders)
        if after:
            self.spacer(self.document, after)

    def columns(self):
        """The sidebar and main columns as one borderless table row."""
        widths = [self.geometry.sidebar_width, self.geometry.main_width]
        if self.geometry.gap:
            widths.insert(1, self.geometry.gap)
        table = self.document.add_table(rows=1, cols=len(widths))
        table.autofit = False
        props = table._tbl.tblPr
        props.find(qn('w:tblW')).set(qn('w:w'), str(sum(widths)))
        props.find(qn('w:tblW')).set(qn('w:type'), 'dxa')
        props.append(element('w:tblInd', w=0, type='dxa'))
        borders = element('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(element('w:' + side, val='nil'))
        props.append(borders)
        for index, width in enumerate(widths):
            table.columns[index].width = Twips(width)
            cell = table.cell(0, index)
            cell.width = Twips(width)
            self.clear_cell(cell)
            region = self.request.document.regions[1] if index == 0 else self.request.document.regions[2] if index == len(widths) - 1 else None
            if region is not None:
                self.blocks(cell, region.blocks)
                # A column's first paragraph opens the body; keeping it with the next one adds nothing there,
                # but converters then move the whole column row to the next page and leave page one empty.
                keep = cell.paragraphs[0]._p.pPr.find(qn('w:keepNext')) if cell.paragraphs and cell.paragraphs[0]._p.pPr is not None else None
                if keep is not None:
                    keep.set(qn('w:val'), '0')
            if not cell.paragraphs:
                self.spacer(cell, 1)

    def render(self, workspace):
        growth = self.header()
        self.footer()
        self.first_page_spacer(growth)
        if self.geometry.sidebar_width:
            self.columns()
        else:
            # One column: the body itself is the main column, in reading order, with no layout table.
            for region in self.request.document.regions[1:]:
                self.blocks(self.document, region.blocks)
        logo = self.recipe.style.header.logo
        if logo is not None:
            asset = next(asset for asset in self.request.package.assets if asset.id == logo.asset)
            self.logo = decode_logo(asset)
        paths = CapturedAssets({selection.font.asset_id: selection.font for selection in self.selections.values()}, self.logo,
            self.icons).materialize(workspace)
        for paragraph, icon, width, height, x, y, right in self.placed:
            self.floating_image(paragraph, paths[icon], width, height, x, y, label='', horizontal='column',
                vertical='paragraph', right=right)
        if logo is not None:
            ratio = min(logo.width / self.logo.width, logo.height / self.logo.height)
            width, height = self.logo.width * ratio, self.logo.height * ratio
            x, top = logo.x + (logo.width - width) / 2, logo.top + (logo.height - height) / 2
            if x + width > self.recipe.style.page.width or top + height > self.recipe.style.page.height:
                raise RenderFailure('RICH_EXPORT_UNSUPPORTED')
            self.floating_image(self.document.sections[0].first_page_header.paragraphs[0], paths[logo.asset], width, height, x, top,
                label='Captured template logo')
        raw = BytesIO()
        self.order_properties()
        self.document.save(raw)
        if len(raw.getvalue()) > DOCX_BYTES:
            raise RenderFailure('RENDER_OUTPUT_LIMIT')
        return embed_captured_fonts(raw.getvalue(), list(self.selections.values()))

    def order_properties(self):
        names = ('pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
            'numPr', 'pBdr', 'shd', 'tabs', 'spacing', 'ind', 'jc', 'outlineLvl', 'rPr')
        order = {qn('w:' + name): index for index, name in enumerate(names)}
        for part in self.document.part.package.parts:
            root = getattr(part, 'element', None)
            if root is None:
                continue
            for props in root.iter(qn('w:pPr')):
                props[:] = sorted(props, key=lambda child: order.get(child.tag, len(order)))
            table_names = ('tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
                'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd',
                'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription', 'tblPrChange')
            table_order = {qn('w:' + name): index for index, name in enumerate(table_names)}
            for props in root.iter(qn('w:tblPr')):
                props[:] = sorted(props, key=lambda child: table_order.get(child.tag, len(table_order)))


def located(block):
    """What an export refusal says about its place: the block, and the text the user can find it by."""
    if block is None:
        return {}
    spans = authored_spans(block) if isinstance(block.content, list) else ()
    return {'blockId': block.id, 'fragment': ' '.join(''.join(getattr(span, 'text', None) or '' for span, _link in spans).split())[:60]}


def render_captured_docx(request, workspace):
    request = TemplateRenderRequestV1.model_validate(request.model_dump(mode='json', by_alias=True))
    if request.renderer != renderer_identity():
        raise RenderFailure('RENDERER_MISMATCH')
    writer = CapturedWriter(request)
    try:
        writer.prepare()
        return writer.render(workspace)
    except RenderFailure as error:
        raise RenderFailure(error.code, located(writer.block)) from None
