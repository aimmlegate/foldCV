"""Flowing editable DOCX from the authoritative rich tree."""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
import uuid
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt, RGBColor
from lxml import etree

from ..models.rich_document import Hyperlink, RichDocumentV1
from ..rich_limits import DOCX_BYTES
from .rich_template import PACKAGE_SHA256, ROOT, require_rich_template
from .template import RenderFailure

RENDERER_VERSION = "1.0.0"
INITIAL_BLOCKS = {"paragraph", "heading", "section", "bullet_list", "numbered_list", "check_list", "toggle_list", "quote", "code", "divider"}


def renderer_sha256():
    return hashlib.sha256(Path(__file__).read_bytes()).hexdigest()


def validate_export(document: RichDocumentV1):
    if any(block.type not in INITIAL_BLOCKS for block in document.blocks()):
        raise RenderFailure("RICH_EXPORT_UNSUPPORTED")


def current_snapshot(workspace, artifacts):
    if workspace.project is None:
        return None
    validate_export(workspace.project.document)
    if workspace.project.template != require_rich_template().manifest:
        raise RenderFailure("RICH_TEMPLATE_MISMATCH")
    for snapshot in reversed(workspace.project.renders):
        blob = artifacts.get(snapshot.id)
        if (snapshot.document_revision == workspace.document_revision
            and snapshot.facts_revision == workspace.project.facts.revision
            and snapshot.template_sha256 == PACKAGE_SHA256
            and snapshot.renderer_version == RENDERER_VERSION
            and snapshot.renderer_sha256 == renderer_sha256()
            and blob is not None and len(blob) == snapshot.docx_size <= DOCX_BYTES
            and hashlib.sha256(blob).hexdigest() == snapshot.docx_sha256):
            return snapshot, blob
    return None


def element(tag, **attributes):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn("w:" + key), str(value))
    return node


def fill(parent, color):
    if color != "default":
        parent.append(element("w:shd", fill=color.removeprefix("#"), val="clear"))


def embed_fonts(raw, spec):
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                  "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    relationship_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    type_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    with ZipFile(BytesIO(raw)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    table = etree.fromstring(parts["word/fontTable.xml"])
    relationships = etree.Element(f"{{{relationship_ns}}}Relationships")
    for index, face in enumerate(spec["fonts"]):
        name = face["wordFamily"]
        found = table.xpath("w:font[@w:name=$name]", namespaces=namespaces, name=name)
        font = found[0] if found else element("w:font", name=name)
        if not found:
            table.append(font)
        identifier = f"rIdFoldCV{index}"
        key = uuid.UUID(hex=face["sha256"][:32])
        mask = key.bytes[::-1]
        content = bytearray((ROOT / face["path"]).read_bytes())
        for byte in range(32):
            content[byte] ^= mask[byte % 16]
        path = f"fonts/{face['face']}.odttf"
        parts["word/" + path] = bytes(content)
        embedded = element("w:embedBold" if face["wordBold"] else "w:embedRegular", fontKey="{" + str(key).upper() + "}")
        embedded.set(qn("r:id"), identifier)
        font.append(embedded)
        etree.SubElement(relationships, f"{{{relationship_ns}}}Relationship", Id=identifier,
            Type=RELATIONSHIP_TYPE.FONT, Target=path)
    settings = etree.fromstring(parts["word/settings.xml"])
    settings.append(element("w:embedTrueTypeFonts"))
    settings.append(element("w:saveSubsetFonts", val="0"))
    types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(types, f"{{{type_ns}}}Default", Extension="odttf",
        ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont")
    if not types.xpath("*[local-name()='Default'][@Extension='txt']"):
        etree.SubElement(types, f"{{{type_ns}}}Default", Extension="txt", ContentType="text/plain")
    parts["word/fonts/LICENSE.txt"] = (ROOT / spec["fontEmbedding"]["license"]).read_bytes()
    for path, tree in (("word/fontTable.xml", table), ("word/_rels/fontTable.xml.rels", relationships),
                       ("word/settings.xml", settings), ("[Content_Types].xml", types)):
        parts[path] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts["docProps/core.xml"] = b'<?xml version="1.0"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"/>'
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for path, blob in parts.items():
            archive.writestr(path, blob)
    return output.getvalue()


class Writer:
    def __init__(self, spec):
        self.spec = spec
        self.fonts = {face["face"]: face for face in spec["fonts"]}
        self.document = Document()
        self.numbering = self.document.part.numbering_part.element
        self.next_numbering = 6000
        for identifier, bullet in ((5000, True), (5001, False)):
            abstract = element("w:abstractNum", abstractNumId=identifier)
            abstract.append(element("w:multiLevelType", val="multilevel"))
            for level in range(9):
                item = element("w:lvl", ilvl=level)
                item.append(element("w:start", val=1))
                item.append(element("w:numFmt", val="bullet" if bullet else "decimal"))
                item.append(element("w:lvlText", val="•" if bullet else f"%{level + 1}."))
                item.append(element("w:lvlJc", val="left"))
                props = element("w:pPr")
                props.append(element("w:ind", left=720 + level * 360, hanging=360))
                item.append(props)
                run = element("w:rPr")
                run.append(element("w:rFonts", ascii="Helvetica" if bullet else "Source Sans Pro",
                    hAnsi="Helvetica" if bullet else "Source Sans Pro"))
                run.append(element("w:sz", val=28 if bullet else 20))
                item.append(run); abstract.append(item)
            self.numbering.append(abstract)
        self.bullet_numbering = self.new_numbering(5000)
        self.document.core_properties.author = ""
        self.document.core_properties.last_modified_by = ""
        section = self.document.sections[0]
        section.page_width, section.page_height = Pt(spec["page"]["width"]), Pt(spec["page"]["height"])
        section.top_margin, section.bottom_margin = Pt(14), Pt(72)
        section.left_margin = Pt(spec["regions"]["sidebar"]["x"])
        section.right_margin = Pt(spec["page"]["width"] - spec["regions"]["main"]["x"] - spec["regions"]["main"]["width"])
        normal = self.document.styles["Normal"]
        normal.font.name = self.fonts["light"]["wordFamily"]
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.line_spacing = Pt(12.57)
        normal.paragraph_format.keep_together = False
        normal.paragraph_format.keep_with_next = False

    def run(self, paragraph, text, styles, face, size):
        if styles.bold:
            face = "bold"
        font = self.fonts[face]
        run = paragraph.add_run(text)
        run.font.name = "Courier New" if styles.code else font["wordFamily"]
        run.font.size = Pt(size)
        run.bold, run.italic, run.underline = font["wordBold"], styles.italic, styles.underline
        run.font.strike = styles.strike
        fonts = run._r.get_or_add_rPr().rFonts
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attribute), run.font.name)
        run.font.color.rgb = RGBColor.from_string(styles.text_color[1:] if styles.text_color != "default" else "171717")
        fill(run._r.get_or_add_rPr(), styles.background_color)
        return run

    def inline(self, paragraph, content, face, size, text_color="default"):
        def effective(styles):
            return styles.model_copy(update={"text_color": text_color}) if styles.text_color == "default" else styles
        for span in content:
            if isinstance(span, Hyperlink):
                link = element("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(span.href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                for child in span.content:
                    run = self.run(paragraph, child.text, effective(child.styles), face, size)
                    link.append(run._r)
                paragraph._p.append(link)
            else:
                self.run(paragraph, span.text, effective(span.styles), face, size)

    def new_numbering(self, abstract_id, start=1, level=0):
        identifier = self.next_numbering; self.next_numbering += 1
        node = element("w:num", numId=identifier)
        node.append(element("w:abstractNumId", val=abstract_id))
        override = element("w:lvlOverride", ilvl=level)
        override.append(element("w:startOverride", val=start))
        node.append(override); self.numbering.append(node)
        return identifier

    def blocks(self, container, blocks, width, region, depth=0, focus=False):
        sequence = {"numbering": None, "index": 0, "previous": None}
        for block in blocks:
            self.block(container, block, width, region, depth, focus, sequence)

    def block(self, container, block, width, region, depth=0, focus=False, sequence=None):
        paragraph = container.add_paragraph()
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.keep_with_next = False
        # Word's character wrapping preserves long tokens rather than inserting characters.
        paragraph._p.get_or_add_pPr().append(element("w:wordWrap", val="0"))
        paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[getattr(block.props, "alignment", "left")]
        fill(paragraph._p.get_or_add_pPr(), getattr(block.props, "background_color", "default"))
        if depth:
            paragraph.paragraph_format.left_indent = Pt(min(min(depth, 12) * 18, max(0, width - 24)))
        state = sequence if sequence is not None else {"numbering": None, "index": 0, "previous": None}
        numbered = block.type == "numbered_list"
        if numbered:
            start = block.props.start
            if state["previous"] != "numbered_list" or start is not None:
                state["numbering"] = self.new_numbering(5001, start or 1, min(depth, 8))
                state["index"] = start or 1
            else:
                state["index"] += 1
        else:
            state["numbering"] = None
        state["previous"] = block.type
        if block.type in {"bullet_list", "numbered_list"} and depth <= 8:
            numbering = element("w:numPr")
            numbering.append(element("w:ilvl", val=depth))
            numbering.append(element("w:numId", val=state["numbering"] if numbered else self.bullet_numbering))
            paragraph._p.get_or_add_pPr().append(numbering)
            indent = min(36 + depth * 18, max(18, width - 24))
            paragraph.paragraph_format.left_indent = Pt(indent)
            paragraph.paragraph_format.first_line_indent = Pt(-18)
        elif block.type in {"bullet_list", "numbered_list", "check_list", "toggle_list"}:
            from ..models.rich_document import InlineStyles
            marker = f"{state['index']}. " if numbered else "☑ " if block.type == "check_list" and block.props.checked else "☐ " if block.type == "check_list" else "▾ " if block.type == "toggle_list" else "• "
            self.run(paragraph, marker, InlineStyles(), "regular", 10)
            paragraph.paragraph_format.left_indent = Pt(min(36 + min(depth, 12) * 18, max(18, width - 24)))
            paragraph.paragraph_format.first_line_indent = Pt(-18)
        face, size = "light", 10
        if block.type == "heading":
            paragraph.style = self.document.styles[f"Heading {block.props.level}"]
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.keep_together = False
        if region == "header" and block.type == "heading" and block.props.level in {1, 2}:
            face, size = ("black", 18) if block.type == "heading" and block.props.level == 1 else ("semibold", 16)
        elif block.type == "heading":
            face, size = ("bold", 14) if block.props.level == 1 else ("semibold", 12 if block.props.level == 2 else 10)
        elif block.type == "section":
            face, size = "bold", 10 if region == "sidebar" else 14
        elif focus:
            face, size = "semibold", 10
            paragraph.paragraph_format.space_after = Pt(25.43)
        if block.type == "quote":
            paragraph.paragraph_format.left_indent = Pt(min(12 + min(depth, 12) * 18, max(0, width - 24)))
            border = element("w:pBdr"); border.append(element("w:left", val="single", sz=6, color="888888", space=4))
            paragraph._p.get_or_add_pPr().append(border)
        if block.type == "divider":
            border = element("w:pBdr"); border.append(element("w:bottom", val="single", sz=6, color="888888"))
            paragraph._p.get_or_add_pPr().append(border)
        elif block.type == "code":
            from ..models.rich_document import InlineStyles
            self.run(paragraph, block.content, InlineStyles(code=True), "regular", 10)
        else:
            self.inline(paragraph, block.content, face, size, getattr(block.props, "text_color", "default"))
        self.blocks(container, block.children, width, region, depth if block.type == "section" else depth + 1,
            block.type == "section" and block.props.presentation == "focus")

    def render(self, document):
        spec = self.spec
        self.blocks(self.document, document.regions[0].blocks, spec["header"]["textPreferredWidth"], "header")
        table = self.document.add_table(rows=1, cols=3)
        table.autofit = False
        widths = [spec["regions"]["sidebar"]["width"], spec["regions"]["gap"], spec["regions"]["main"]["width"]]
        for index, width in enumerate(widths):
            table.columns[index].width = Pt(width)
            cell = table.cell(0, index)
            cell.width = Pt(width)
            margins = element("w:tcMar")
            for side in ("top", "start", "bottom", "end"):
                margins.append(element("w:" + side, w="0", type="dxa"))
            cell._tc.get_or_add_tcPr().append(margins)
            empty = cell.paragraphs[0]
            empty._p.getparent().remove(empty._p)
            if index == 1:
                cell.add_paragraph()
        for region, index in zip(document.regions[1:], (0, 2)):
            self.blocks(table.cell(0, index), region.blocks, widths[index], region.role)
            if not region.blocks:
                table.cell(0, index).add_paragraph()
        footer = self.document.sections[0].footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        from ..models.rich_document import InlineStyles
        self.run(footer, spec["footer"]["text"], InlineStyles(), "light", 12)
        output = BytesIO()
        self.document.save(output)
        return embed_fonts(output.getvalue(), spec)


def render_rich_docx(document: RichDocumentV1, images=None):
    validate_export(document)
    content = Writer(require_rich_template().specification).render(document)
    if len(content) > DOCX_BYTES:
        raise RenderFailure("RENDER_OUTPUT_LIMIT")
    return content
